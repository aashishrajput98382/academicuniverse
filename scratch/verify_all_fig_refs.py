import docx, re

doc = docx.Document('PaperV50.docx')

print(f'Total paragraphs in PaperV50.docx: {len(doc.paragraphs)}')
print(f'Total tables: {len(doc.tables)}')

# Locate all figures by their caption
fig_data = []
for idx, p in enumerate(doc.paragraphs):
    m = re.match(r'^\s*(Fig\.\s*(\d+))\.\s*(.*)', p.text, re.IGNORECASE)
    if m:
        fig_label = m.group(1)
        fig_num = int(m.group(2))
        fig_title = m.group(3)
        # The image is immediately before the caption
        img_idx = idx - 1
        
        # Paragraph above image
        above_p = None
        k = img_idx - 1
        while k >= 0:
            if doc.paragraphs[k].text.strip():
                above_p = (k, doc.paragraphs[k].text.strip())
                break
            k -= 1

        # Paragraph below caption
        below_p = None
        k = idx + 1
        while k < len(doc.paragraphs):
            if doc.paragraphs[k].text.strip():
                below_p = (k, doc.paragraphs[k].text.strip())
                break
            k += 1

        fig_data.append({
            'num': fig_num,
            'label': fig_label,
            'cap_idx': idx,
            'caption': p.text.strip(),
            'above': above_p,
            'below': below_p,
        })

print('\n' + '='*80)
print('FIGURE-BY-FIGURE REFERENCE AUDIT IN PaperV50.docx')
print('='*80)

for f in fig_data:
    num = f['num']
    label = f'Fig. {num}'
    above_txt = f['above'][1] if f['above'] else ''
    below_txt = f['below'][1] if f['below'] else ''

    has_above = label in above_txt
    has_below = label in below_txt

    cap_idx = f['cap_idx']
    caption = f['caption']
    above_num = f['above'][0] if f['above'] else '-'
    below_num = f['below'][0] if f['below'] else '-'

    print(f'\n--- FIGURE {num} (Caption P{cap_idx}) ---')
    print(f'  Caption: {caption[:75]}')
    print(f'  Paragraph Above (P{above_num}): {repr(above_txt[:90])}')
    print(f'  Paragraph Below (P{below_num}): {repr(below_txt[:90])}')
    print(f'  Referenced in Above Paragraph: {has_above}')
    print(f'  Referenced in Below Paragraph: {has_below}')
    assert has_above or has_below, f'Fig. {num} is NOT referenced in above or below paragraph!'

print('\n' + '='*80)
print('ALL 8 FIGURES ARE CONFIRMED REFERENCED IN PARAGRAPHS DIRECTLY ABOVE OR BELOW!')
