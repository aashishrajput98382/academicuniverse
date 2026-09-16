import docx
import xml.etree.ElementTree as ET

doc = docx.Document('PaperV50.docx')

for i, p in enumerate(doc.paragraphs):
    if p.text.strip().startswith('Fig.') or p.text.strip().startswith('Figure'):
        print(f"Caption P{i}: {p.text.strip()}")
        # Check preceding paragraph for blip
        if i > 0:
            xml = doc.paragraphs[i-1]._element.xml
            if 'blip' in xml:
                root = ET.fromstring(xml)
                for elem in root.iter():
                    if elem.tag.endswith('blip'):
                        rId = elem.attrib.get('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed')
                        part = doc.part.related_parts[rId]
                        print(f"  Preceding P{i-1} has image: {part.partname} (rId: {rId})")
        # Check current paragraph
        xml = p._element.xml
        if 'blip' in xml:
            root = ET.fromstring(xml)
            for elem in root.iter():
                if elem.tag.endswith('blip'):
                    rId = elem.attrib.get('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed')
                    part = doc.part.related_parts[rId]
                    print(f"  Current P{i} has image: {part.partname} (rId: {rId})")
