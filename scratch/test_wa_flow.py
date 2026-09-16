import docx
from docx.shared import Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
import pymupdf
import win32com.client, os, shutil

test_docx = r'c:\github\academicuniverse\scratch\test_wa_flow.docx'
test_pdf = r'c:\github\academicuniverse\scratch\test_wa_flow.pdf'
src_backup = r'c:\github\academicuniverse\ashmitxnaik\Deepfake_Video_Authentication_Updated_Research_Paper_With_Flowchart (1)_backup.docx'
wa_img = r'c:\github\academicuniverse\ashmitxnaik\WhatsApp Image 2026-09-16 at 2.51.05 PM.jpeg'

shutil.copyfile(src_backup, test_docx)

doc = docx.Document(test_docx)

# Locate target block
t_start = None
for i, p in enumerate(doc.paragraphs):
    if 'Input Video' in p.text and 'Facial Sequence' in p.text:
        t_start = i
        break

to_rem = []
for j in range(t_start + 1, t_start + 15):
    pj = doc.paragraphs[j]
    to_rem.append(pj)
    if 'Authentication Dashboard' in pj.text:
        break

# Replace target start with image
p_img = doc.paragraphs[t_start]
p_img.text = ""
p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p_img.add_run()
r.add_picture(wa_img, width=Inches(3.48))

# Caption
p_cap = to_rem[0]
p_cap.text = "Fig. 1. System workflow and operational architecture."
p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
for run in p_cap.runs:
    run.font.size = 101600 # 8 pt
    run.italic = True

for pj in to_rem[1:]:
    pe = pj._p
    pe.getparent().remove(pe)

# 1. Remove column break at Batch size
for p in doc.paragraphs:
    if 'Batch size' in p.text:
        col_brs = p._p.xpath('.//w:br[@w:type="column"]')
        for b in col_brs:
            b.getparent().remove(b)

# 2. Make sections continuous so text flows naturally across columns
for s_idx in [9, 10]:
    if s_idx < len(doc.sections):
        s = doc.sections[s_idx]
        s.start_type = docx.enum.section.WD_SECTION_START.CONTINUOUS
        type_els = s._sectPr.xpath('./w:type')
        if type_els:
            type_els[0].set(docx.oxml.ns.qn('w:val'), 'continuous')
        else:
            t = docx.oxml.OxmlElement('w:type')
            t.set(docx.oxml.ns.qn('w:val'), 'continuous')
            s._sectPr.insert(0, t)

doc.save(test_docx)

# Export PDF
word = win32com.client.Dispatch('Word.Application')
word.Visible = False
d = word.Documents.Open(os.path.abspath(test_docx))
d.SaveAs(os.path.abspath(test_pdf), FileFormat=17)
d.Close()
word.Quit()

pdf = pymupdf.open(test_pdf)
print(f"Total pages: {len(pdf)}")
for pno in range(len(pdf)):
    pix = pdf[pno].get_pixmap(dpi=150)
    pix.save(f"scratch/wa_flow_page_{pno+1}.png")
    print(f"Rendered scratch/wa_flow_page_{pno+1}.png")
