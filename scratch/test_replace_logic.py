import docx
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
import shutil

src = r'c:\github\academicuniverse\ashmitxnaik\Deepfake_Video_Authentication_Updated_Research_Paper_With_Flowchart (1).docx'
dst = r'c:\github\academicuniverse\scratch\test_replaced.docx'
shutil.copyfile(src, dst)

doc = docx.Document(dst)

# Find the start paragraph (paragraph 96 with 'Input Video')
target_start = None
for i, p in enumerate(doc.paragraphs):
    if 'Input Video' in p.text and 'Facial Sequence' in p.text:
        target_start = i
        break

print(f"Target start index: {target_start}")
if target_start is not None:
    # Find all paragraphs in that block up to 'Authentication Dashboard'
    to_remove = []
    for j in range(target_start + 1, target_start + 15):
        pj = doc.paragraphs[j]
        to_remove.append(pj)
        if 'Authentication Dashboard' in pj.text:
            break
    
    print(f"Found {len(to_remove)} paragraphs to remove:")
    for pj in to_remove:
        print(f"  - {repr(pj.text)}")
    
    # Replace target_start paragraph content
    p_img = doc.paragraphs[target_start]
    p_img.text = ""
    p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p_img.add_run()
    img_path = r'c:\github\academicuniverse\ashmitxnaik\WhatsApp Image 2026-09-16 at 2.51.05 PM.jpeg'
    run.add_picture(img_path, width=Inches(3.6))
    
    # Remove following paragraphs in that block
    for pj in to_remove:
        p_elem = pj._p
        p_elem.getparent().remove(p_elem)
        
    doc.save(dst)
    print("Successfully saved test_replaced.docx!")
