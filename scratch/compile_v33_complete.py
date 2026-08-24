import os
import re
import subprocess
import win32com.client
import docx
from docx.shared import Pt, RGBColor
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls
from pathlib import Path

workspace = Path(__file__).resolve().parents[1]
v32_docx_path = workspace / "docs" / "paper" / "PaperV32_Ollama_Primary.docx"
v33_docx_path = workspace / "docs" / "paper" / "PaperV33_Ollama_Primary.docx"
v33_pdf_path = workspace / "docs" / "paper" / "PaperV33_Ollama_Primary.pdf"

doc = docx.Document(v32_docx_path)

# Append Appendix A
p_app_a = doc.add_paragraph()
p_app_a.paragraph_format.space_before = Pt(14)
p_app_a.paragraph_format.space_after = Pt(4)
r_a = p_app_a.add_run("APPENDIX A: CANONICAL NORMALIZATION RULES & ALIAS MAPPINGS")
r_a.font.name = "Times New Roman"
r_a.font.size = Pt(10.0)
r_a.font.bold = True
r_a.font.color.rgb = RGBColor(27, 54, 93)

app_a_items = [
    "To facilitate reproduction and institutional deployment, the formal deterministic transformation rules implemented in the Six-Stage Semantic CanonicalNormalizer (N) are specified below:",
    "1. Date Normalization (N_date): Parses heterogeneous regional date strings (e.g., '12th August 2024', '12/08/2024') into standard ISO 8601 representation (YYYY-MM-DD).",
    "2. Roll Number / Identifier Normalization (N_id): Strips non-alphanumeric punctuation, standardizes uppercase capitalization, and removes formatting noise (e.g., '2023-CS/042' -> '2023CS042').",
    "3. Numeric & Grade Normalization (N_num): Extracts floating-point scalars from mixed textual expressions (e.g., '8.45 / 10.0 CGPA' -> '8.45') with strict two-decimal precision.",
    "4. Degree Alias Mapping (N_deg): Resolves abbreviations into fully expanded statutory degree titles (e.g., 'B.Tech' -> 'Bachelor of Technology').",
    "5. Honorific & Whitespace Normalization (N_ws): Strips leading honorific prefixes ('Mr.', 'Ms.', 'Dr.') and collapses internal redundant whitespace sequences.",
    "6. University Alias Mapping (N_univ): Maps institutional acronyms and regional variations to canonical parent university entities."
]

for item in app_a_items:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(1)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.05
    r = p.add_run(item)
    r.font.name = "Times New Roman"
    r.font.size = Pt(8.0)
    r.font.color.rgb = RGBColor(30, 41, 59)

# Append Appendix B
p_app_b = doc.add_paragraph()
p_app_b.paragraph_format.space_before = Pt(12)
p_app_b.paragraph_format.space_after = Pt(4)
r_b = p_app_b.add_run("APPENDIX B: NINE-CLASS DIAGNOSTIC OCR ERROR TAXONOMY SPECIFICATION")
r_b.font.name = "Times New Roman"
r_b.font.size = Pt(10.0)
r_b.font.bold = True
r_b.font.color.rgb = RGBColor(27, 54, 93)

app_b_items = [
    "The automated diagnostic classifier evaluates field pairs (y, y_hat) against the following sequential, mutually exclusive decision rules:",
    "1. EXACT_MATCH: Unnormalized exact string equality (y == y_hat).",
    "2. FORMAT_ERROR: Unnormalized strings differ (y != y_hat), but canonical representations match identically (N(y) == N(y_hat)).",
    "3. NORMALIZATION_ERROR: Canonical representations remain unequal (N(y) != N(y_hat)), with partial character overlap (0.5 <= LevSim < 1.0).",
    "4. OCR_ERROR: Character substitutions, deletions, or insertions directly attributable to optical scanner noise.",
    "5. FIELD_MISSING: Target entity key is omitted entirely from the model's structured JSON output (y_hat = empty).",
    "6. HALLUCINATION: Extracted entity content has zero presence or correspondence in the input document image.",
    "7. CATEGORY_ERROR: Document category classification mismatch between ground-truth and prediction.",
    "8. PARTIAL_MATCH: Extracted string is a strict sub-phrase or truncated fragment of the reference entity (0 < LevSim < 0.5).",
    "9. LOW_CONFIDENCE: Extracted field prediction falls below the calibrated model confidence cutoff threshold (tau < 0.50)."
]

for item in app_b_items:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(1)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.05
    r = p.add_run(item)
    r.font.name = "Times New Roman"
    r.font.size = Pt(8.0)
    r.font.color.rgb = RGBColor(30, 41, 59)

# Compact formatting for <= 20 pages
for p in doc.paragraphs:
    txt = p.text.strip()
    if txt.startswith("TABLE ") or txt.startswith("Table "):
        p.paragraph_format.space_before = Pt(3)
        p.paragraph_format.space_after = Pt(1)
    elif txt.startswith("Fig. ") or txt.startswith("FIG. "):
        p.paragraph_format.space_before = Pt(1)
        p.paragraph_format.space_after = Pt(3)
    elif re.match(r'^\d+\.\s+', txt):
        p.paragraph_format.space_before = Pt(6)
        p.paragraph_format.space_after = Pt(2)

    for r in p.runs:
        drawings = r._r.xpath('.//w:drawing')
        if drawings:
            for d in drawings:
                exts = d.xpath('.//wp:extent')
                for ext in exts:
                    ext.set('cx', str(int(6.0 * 914400)))
                    ext.set('cy', str(int(3.2 * 914400)))
                a_exts = d.xpath('.//a:ext')
                for a_ext in a_exts:
                    a_ext.set('cx', str(int(6.0 * 914400)))
                    a_ext.set('cy', str(int(3.2 * 914400)))

doc.save(v33_docx_path)
print(f"[SUCCESS] Saved PaperV33_Ollama_Primary.docx with Appendix A & Appendix B!")

# Export PDF
subprocess.run(["taskkill", "/F", "/IM", "WINWORD.EXE"], capture_output=True)

word = win32com.client.DispatchEx("Word.Application")
word.Visible = False
word.DisplayAlerts = False
try:
    doc_word = word.Documents.Open(os.path.abspath(str(v33_docx_path)))
    doc_word.SaveAs(os.path.abspath(str(v33_pdf_path)), FileFormat=17)
    page_count = doc_word.ComputeStatistics(2)
    doc_word.Close(False)
    print(f"[SUCCESS] Exported high-quality PDF: {v33_pdf_path.name} ({page_count} Pages!)")
finally:
    word.Quit()
