import docx
from docx.shared import Pt, Inches
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls
import win32com.client
import os
from pathlib import Path

workspace = Path(__file__).resolve().parents[1]
v22_docx_path = workspace / "docs" / "paper" / "PaperV22_Ollama_Primary.docx"
v22_pdf_path = workspace / "docs" / "paper" / "PaperV22_Ollama_Primary.pdf"

doc = docx.Document(v22_docx_path)

# 1. Compact table cell margins and text spacing across all tables
for tbl in doc.tables:
    for row in tbl.rows:
        for cell in row.cells:
            tcPr = cell._tc.get_or_add_tcPr()
            tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="25" w:type="dxa"/><w:bottom w:w="25" w:type="dxa"/><w:left w:w="35" w:type="dxa"/><w:right w:w="35" w:type="dxa"/></w:tcMar>')
            tcPr.append(tcMar)
            for p in cell.paragraphs:
                p.paragraph_format.space_before = Pt(1)
                p.paragraph_format.space_after = Pt(1)
                p.paragraph_format.line_spacing = 1.0
                for r in p.runs:
                    r.font.name = "Times New Roman"
                    r.font.size = Pt(8.0)

# 2. Compact paragraph spaces around figures and headings
for p in doc.paragraphs:
    txt = p.text.strip()
    if txt.startswith("TABLE ") or txt.startswith("Table "):
        p.paragraph_format.space_before = Pt(6)
        p.paragraph_format.space_after = Pt(2)
    elif txt.startswith("Fig. ") or txt.startswith("FIG. "):
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(6)
    elif txt.startswith("1. ") or txt.startswith("2. ") or txt.startswith("3. ") or txt.startswith("4. ") or txt.startswith("5. ") or txt.startswith("6. ") or txt.startswith("7. "):
        p.paragraph_format.space_before = Pt(10)
        p.paragraph_format.space_after = Pt(4)
    elif txt.startswith("4.") or txt.startswith("5."):
        p.paragraph_format.space_before = Pt(6)
        p.paragraph_format.space_after = Pt(2)

# 3. Scale figures to compact vertical size
for p in doc.paragraphs:
    for r in p.runs:
        drawings = r._r.xpath('.//w:drawing')
        if drawings:
            for d in drawings:
                exts = d.xpath('.//wp:extent')
                for ext in exts:
                    cx = int(ext.get('cx', 0))
                    cy = int(ext.get('cy', 0))
                    if cy > 4.0 * 914400:
                        new_cy = int(4.0 * 914400)
                        new_cx = int(cx * (new_cy / cy))
                        ext.set('cx', str(new_cx))
                        ext.set('cy', str(new_cy))
                a_exts = d.xpath('.//a:ext')
                for a_ext in a_exts:
                    cx = int(a_ext.get('cx', 0))
                    cy = int(a_ext.get('cy', 0))
                    if cy > 4.0 * 914400:
                        new_cy = int(4.0 * 914400)
                        new_cx = int(cx * (new_cy / cy))
                        a_ext.set('cx', str(new_cx))
                        a_ext.set('cy', str(new_cy))

doc.save(v22_docx_path)
print("[SUCCESS] Saved compacted PaperV22_Ollama_Primary.docx!")

# Export PDF and check page count
word = None
try:
    word = win32com.client.Dispatch("Word.Application")
    word.Visible = False
    doc_word = word.Documents.Open(os.path.abspath(str(v22_docx_path)))
    doc_word.SaveAs(os.path.abspath(str(v22_pdf_path)), FileFormat=17)
    page_count = doc_word.ComputeStatistics(2)
    doc_word.Close()
    print(f"[SUCCESS] Exported high-quality PDF: {v22_pdf_path.name} ({page_count} Pages!)")
except Exception as e:
    print(f"Word COM PDF Export Error: {e}")
finally:
    if word:
        try: word.Quit()
        except: pass
