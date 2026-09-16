import docx, re

doc = docx.Document('PaperV50.docx')

print('Checking paragraphs...')
for i, p in enumerate(doc.paragraphs):
    full_text = p.text
    if re.search(r'\b(TABLE|Table)\s*(\d+)\b', full_text):
        # check each run
        matches_in_full = re.findall(r'\b(TABLE|Table)\s*(\d+)\b', full_text)
        matches_in_runs = []
        for r in p.runs:
            matches_in_runs.extend(re.findall(r'\b(TABLE|Table)\s*(\d+)\b', r.text))
        if len(matches_in_full) != len(matches_in_runs):
            print(f'P{i} has split runs! Full: {matches_in_full}, Runs: {matches_in_runs}')
        else:
            print(f'P{i} OK: {matches_in_runs}')

print('Checking tables...')
for t_idx, tbl in enumerate(doc.tables):
    for r_idx, row in enumerate(tbl.rows):
        for c_idx, cell in enumerate(row.cells):
            for p in cell.paragraphs:
                full_text = p.text
                if re.search(r'\b(TABLE|Table)\s*(\d+)\b', full_text):
                    matches_in_full = re.findall(r'\b(TABLE|Table)\s*(\d+)\b', full_text)
                    matches_in_runs = []
                    for r in p.runs:
                        matches_in_runs.extend(re.findall(r'\b(TABLE|Table)\s*(\d+)\b', r.text))
                    if len(matches_in_full) != len(matches_in_runs):
                        print(f'T{t_idx} R{r_idx} C{c_idx} has split runs! Full: {matches_in_full}, Runs: {matches_in_runs}')
                    else:
                        print(f'T{t_idx} R{r_idx} C{c_idx} OK: {matches_in_runs}')
