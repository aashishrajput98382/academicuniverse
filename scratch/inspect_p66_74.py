import docx

doc = docx.Document('PaperV50.docx')
for i in range(66, 75):
    print(f'P{i}: text="{doc.paragraphs[i].text}", runs={[r.text for r in doc.paragraphs[i].runs]}')
