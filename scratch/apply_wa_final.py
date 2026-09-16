import docx
from docx.shared import Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
import win32com.client, os, shutil

src_backup = r'c:\github\academicuniverse\ashmitxnaik\Deepfake_Video_Authentication_Updated_Research_Paper_With_Flowchart (1)_backup.docx'
target_docx = r'c:\github\academicuniverse\ashmitxnaik\Deepfake_Video_Authentication_Updated_Research_Paper_With_Flowchart (1).docx'
target_pdf = r'c:\github\academicuniverse\ashmitxnaik\Deepfake_Video_Authentication_Updated_Research_Paper_With_Flowchart (1).pdf'
wa_img = r'c:\github\academicuniverse\ashmitxnaik\WhatsApp Image 2026-09-16 at 2.51.05 PM.jpeg'

# Copy clean backup to target
shutil.copyfile(src_backup, target_docx)

doc = docx.Document(target_docx)

# Locate target block
t_start = None
for i, p in enumerate(doc.paragraphs):
    if 'Input Video' in p.text and 'Facial Sequence' in p.text:
        t_start = i
        break

if t_start is None:
    raise ValueError("Target paragraph not found!")

to_rem = []
for j in range(t_start + 1, t_start + 15):
    pj = doc.paragraphs[j]
    to_rem.append(pj)
    if 'Authentication Dashboard' in pj.text:
        break

# Replace target start with exact WhatsApp image
p_img = doc.paragraphs[t_start]
p_img.text = ""
p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p_img.add_run()
# height = 2.0 inches preserves the full 8-page layout without causing empty columns on Page 3 and Page 4
r.add_picture(wa_img, height=Inches(2.0))

# Caption
p_cap = to_rem[0]
p_cap.text = "Fig. 1. System workflow and operational architecture."
p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
for run in p_cap.runs:
    run.font.size = 101600 # 8 pt
    run.italic = True

# Remove remaining flowchart text paragraphs
for pj in to_rem[1:]:
    pe = pj._p
    pe.getparent().remove(pe)

doc.save(target_docx)
print("Updated target docx with exact WhatsApp image.")

# Recompile PDF
word = win32com.client.Dispatch('Word.Application')
word.Visible = False
d = word.Documents.Open(os.path.abspath(target_docx))
d.SaveAs(os.path.abspath(target_pdf), FileFormat=17)
d.Close()
try:
    word.Quit()
except:
    pass
print("Compiled target PDF successfully at:", target_pdf)
