import os
import shutil
import re
import docx
from docx.oxml.text.paragraph import CT_P
from docx.oxml.table import CT_Tbl

DOCX_PATH = 'PaperV50.docx'
BACKUP_PATH = 'scratch/PaperV50_pre_t7_removal_backup.docx'

def execute():
    # 1. Backup
    shutil.copy2(DOCX_PATH, BACKUP_PATH)
    print(f"Backed up {DOCX_PATH} to {BACKUP_PATH}")

    doc = docx.Document(DOCX_PATH)
    print(f"Initial state: {len(doc.paragraphs)} paragraphs, {len(doc.tables)} tables")

    # 2. Identify Table 7 and associated paragraphs
    p63 = doc.paragraphs[63]
    p64 = doc.paragraphs[64]
    t7 = doc.tables[7]
    p65 = doc.paragraphs[65]

    print(f"P63 text: {p63.text[:70]}...")
    print(f"P64 text: {p64.text[:70]}...")
    print(f"T7 dimensions: {len(t7.rows)} rows x {len(t7.columns)} cols")
    print(f"P65 text: {p65.text[:70]}...")

    assert "dry-run" in p63.text.lower() and "table 7" in p63.text.lower(), "P63 mismatch!"
    assert "TABLE 7" in p64.text, "P64 mismatch!"
    assert "dry-run" in p65.text.lower(), "P65 mismatch!"

    # Remove from XML
    p63._element.getparent().remove(p63._element)
    p64._element.getparent().remove(p64._element)
    t7._element.getparent().remove(t7._element)
    p65._element.getparent().remove(p65._element)

    print(f"After removal: {len(doc.paragraphs)} paragraphs, {len(doc.tables)} tables")
    assert len(doc.tables) == 14, f"Expected 14 tables, got {len(doc.tables)}"

    # 3. Renumbering function (non-cascading single pass)
    def renumber_tables(text):
        def repl(match):
            prefix = match.group(1) # 'Table ' or 'TABLE '
            num = int(match.group(2))
            if 8 <= num <= 14:
                return f"{prefix}{num - 1}"
            return match.group(0)
        return re.sub(r'\b(TABLE\s*|Table\s*)(\d+)\b', repl, text)

    # Apply to all paragraphs and their runs
    p_changes = 0
    for p_idx, p in enumerate(doc.paragraphs):
        old_text = p.text
        new_text = renumber_tables(old_text)
        if old_text != new_text:
            p_changes += 1
            print(f"\nUpdating P{p_idx}:")
            print(f"  OLD: {old_text[:90]}...")
            print(f"  NEW: {new_text[:90]}...")
            # Update within runs
            for r in p.runs:
                r_old = r.text
                r_new = renumber_tables(r_old)
                if r_old != r_new:
                    r.text = r_new

    # Apply to all table cells just in case
    tbl_cell_changes = 0
    for t_idx, tbl in enumerate(doc.tables):
        for r_idx, row in enumerate(tbl.rows):
            for c_idx, cell in enumerate(row.cells):
                for p in cell.paragraphs:
                    old_text = p.text
                    new_text = renumber_tables(old_text)
                    if old_text != new_text:
                        tbl_cell_changes += 1
                        for r in p.runs:
                            r.text = renumber_tables(r.text)

    print(f"\nTotal paragraphs updated: {p_changes}")
    print(f"Total table cells updated: {tbl_cell_changes}")

    # 4. Save
    doc.save(DOCX_PATH)
    print(f"Successfully saved changes to {DOCX_PATH}!")

if __name__ == '__main__':
    execute()
