import os
import re
import subprocess
import win32com.client
import docx
from docx.shared import Pt, RGBColor
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls
from pathlib import Path

workspace = Path(__file__).resolve().parents[1]
v30_docx_path = workspace / "docs" / "paper" / "PaperV30_Ollama_Primary.docx"
v32_docx_path = workspace / "docs" / "paper" / "PaperV32_Ollama_Primary.docx"
v32_pdf_path = workspace / "docs" / "paper" / "PaperV32_Ollama_Primary.pdf"

doc = docx.Document(v30_docx_path)

# Load strict references
refs_path = workspace / "scratch" / "formatted_references.txt"
strict_refs = [line.strip() for line in refs_path.read_text(encoding="utf-8").split("\n\n") if line.strip()]
print(f"Loaded {len(strict_refs)} strict references.")

# Find REFERENCES heading in doc
ref_heading_idx = None
for idx, p in enumerate(doc.paragraphs):
    if p.text.strip().upper() == "REFERENCES":
        ref_heading_idx = idx
        break

print(f"Found REFERENCES heading at paragraph index {ref_heading_idx}")

# Format REFERENCES heading
ref_p = doc.paragraphs[ref_heading_idx]
ref_p.text = "REFERENCES"
ref_p.paragraph_format.space_before = Pt(12)
ref_p.paragraph_format.space_after = Pt(4)
for r in ref_p.runs:
    r.font.name = "Times New Roman"
    r.font.size = Pt(10.0)
    r.font.bold = True
    r.font.color.rgb = RGBColor(27, 54, 93) # Navy Blue

# Update subsequent 50 paragraphs with strict references
for i, ref_text in enumerate(strict_refs):
    p_idx = ref_heading_idx + 1 + i
    if p_idx < len(doc.paragraphs):
        p = doc.paragraphs[p_idx]
        p.text = ref_text
    else:
        p = doc.add_paragraph(ref_text)
    
    p.paragraph_format.space_before = Pt(1)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.05
    for r in p.runs:
        r.font.name = "Times New Roman"
        r.font.size = Pt(8.0)
        r.font.color.rgb = RGBColor(30, 41, 59)

# Clean up any trailing paragraphs beyond index (ref_heading_idx + 50)
while len(doc.paragraphs) > ref_heading_idx + 1 + len(strict_refs):
    p_to_del = doc.paragraphs[-1]
    p_to_del._element.getparent().remove(p_to_del._element)

print(f"Final paragraph count: {len(doc.paragraphs)}")
print(f"Last paragraph: {repr(doc.paragraphs[-1].text[:60])}")

# Re-apply table styling across all 14 tables
def set_cell_shading(cell, color_hex):
    shading_elm = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{color_hex}"/>')
    cell._tc.get_or_add_tcPr().append(shading_elm)

for tbl in doc.tables:
    for r_idx, row in enumerate(tbl.rows):
        is_header = (r_idx == 0)
        bg_color = "1B365D" if is_header else ("F7FAFC" if r_idx % 2 == 1 else "FFFFFF")
        for cell in row.cells:
            tcPr = cell._tc.get_or_add_tcPr()
            tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="28" w:type="dxa"/><w:bottom w:w="28" w:type="dxa"/><w:left w:w="35" w:type="dxa"/><w:right w:w="35" w:type="dxa"/></w:tcMar>')
            tcPr.append(tcMar)
            set_cell_shading(cell, bg_color)
            for p in cell.paragraphs:
                p.paragraph_format.space_before = Pt(1)
                p.paragraph_format.space_after = Pt(1)
                p.paragraph_format.line_spacing = 1.0
                for r in p.runs:
                    r.font.name = "Times New Roman"
                    if is_header:
                        r.font.size = Pt(8.0)
                        r.font.bold = True
                        r.font.color.rgb = RGBColor(255, 255, 255)
                    else:
                        r.font.size = Pt(7.2)
                        r.font.color.rgb = RGBColor(30, 41, 59)

# Compact formatting for <= 20 pages
for p in doc.paragraphs:
    txt = p.text.strip()
    if txt.startswith("TABLE ") or txt.startswith("Table "):
        p.paragraph_format.space_before = Pt(3)
        p.paragraph_format.space_after = Pt(1)
    elif txt.startswith("Fig. ") or txt.startswith("FIG. "):
        p.paragraph_format.space_before = Pt(1)
        p.paragraph_format.space_after = Pt(3)
    elif re.match(r'^\d+\.\s+', txt):
        p.paragraph_format.space_before = Pt(6)
        p.paragraph_format.space_after = Pt(2)

    for r in p.runs:
        drawings = r._r.xpath('.//w:drawing')
        if drawings:
            for d in drawings:
                exts = d.xpath('.//wp:extent')
                for ext in exts:
                    ext.set('cx', str(int(6.0 * 914400)))
                    ext.set('cy', str(int(3.2 * 914400)))
                a_exts = d.xpath('.//a:ext')
                for a_ext in a_exts:
                    a_ext.set('cx', str(int(6.0 * 914400)))
                    a_ext.set('cy', str(int(3.2 * 914400)))

doc.save(v32_docx_path)
print(f"[SUCCESS] Saved PaperV32_Ollama_Primary.docx with correct REFERENCES heading position!")

# Export PDF
subprocess.run(["taskkill", "/F", "/IM", "WINWORD.EXE"], capture_output=True)

word = win32com.client.DispatchEx("Word.Application")
word.Visible = False
word.DisplayAlerts = False
try:
    doc_word = word.Documents.Open(os.path.abspath(str(v32_docx_path)))
    doc_word.SaveAs(os.path.abspath(str(v32_pdf_path)), FileFormat=17)
    page_count = doc_word.ComputeStatistics(2)
    doc_word.Close(False)
    print(f"[SUCCESS] Exported high-quality PDF: {v32_pdf_path.name} ({page_count} Pages!)")
finally:
    word.Quit()
