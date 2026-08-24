import docx

doc = docx.Document('docs/paper/PaperV27_Ollama_Primary.docx')
print(f"Total tables: {len(doc.tables)}")
for i, t in enumerate(doc.tables):
    headers = [c.text.strip().replace('\n', ' ') for c in t.rows[0].cells]
    print(f"Table index {i}: {len(t.rows)} rows, {len(t.columns)} cols")
    print(f"  Headers: {headers}")
