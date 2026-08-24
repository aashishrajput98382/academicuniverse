import docx

doc = docx.Document("docs/paper/PaperV31_Ollama_Primary.docx")
print(f"Total paragraphs: {len(doc.paragraphs)}")
for idx in range(max(0, len(doc.paragraphs)-20), len(doc.paragraphs)):
    p = doc.paragraphs[idx]
    print(f"[{idx}] {repr(p.text[:80])}")
