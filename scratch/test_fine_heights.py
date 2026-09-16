import docx
from docx.shared import Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
import pymupdf
import win32com.client, os, shutil

src_backup = r'c:\github\academicuniverse\ashmitxnaik\Deepfake_Video_Authentication_Updated_Research_Paper_With_Flowchart (1)_backup.docx'
wa_img = r'c:\github\academicuniverse\ashmitxnaik\WhatsApp Image 2026-09-16 at 2.51.05 PM.jpeg'

for test_h in [1.6, 1.7, 1.8, 1.9, 2.0]:
    test_docx = rf'c:\github\academicuniverse\scratch\test_wa_h_{test_h}.docx'
    test_pdf = rf'c:\github\academicuniverse\scratch\test_wa_h_{test_h}.pdf'
    shutil.copyfile(src_backup, test_docx)
    
    doc = docx.Document(test_docx)
    t_start = None
    for i, p in enumerate(doc.paragraphs):
        if 'Input Video' in p.text and 'Facial Sequence' in p.text:
            t_start = i
            break
            
    to_rem = []
    for j in range(t_start + 1, t_start + 15):
        pj = doc.paragraphs[j]
        to_rem.append(pj)
        if 'Authentication Dashboard' in pj.text:
            break
            
    p_img = doc.paragraphs[t_start]
    p_img.text = ""
    p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p_img.add_run()
    r.add_picture(wa_img, height=Inches(test_h))
    
    p_cap = to_rem[0]
    p_cap.text = "Fig. 1. Overall architecture and operational workflow."
    p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in p_cap.runs:
        run.font.size = 101600 # 8 pt
        run.italic = True
        
    for pj in to_rem[1:]:
        pe = pj._p
        pe.getparent().remove(pe)
        
    doc.save(test_docx)
    
    word = win32com.client.Dispatch('Word.Application')
    word.Visible = False
    d = word.Documents.Open(os.path.abspath(test_docx))
    d.SaveAs(os.path.abspath(test_pdf), FileFormat=17)
    d.Close()
    try:
        word.Quit()
    except:
        pass
        
    pdf = pymupdf.open(test_pdf)
    print(f"Height {test_h} in -> Total pages: {len(pdf)}")
    p3_text = pdf[2].get_text()
    has_ix = 'IX. TEMPORAL FEATURE LEARNING' in p3_text
    print(f"  P3 has IX. TEMPORAL FEATURE LEARNING: {has_ix}")
