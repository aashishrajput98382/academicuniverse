import docx
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

def update_conclusion_and_future_work(docx_path):
    doc = docx.Document(docx_path)

    concl_p = None
    fw_p = None

    for i, p in enumerate(doc.paragraphs):
        t = p.text.strip()
        if t == '6 CONCLUSION':
            concl_p = doc.paragraphs[i+1]
        elif t == '7 FUTURE WORK':
            fw_p = doc.paragraphs[i+1]

    assert concl_p is not None, 'Conclusion body not found!'
    assert fw_p is not None, 'Future work body not found!'

    concl_text = (
        "Benchmarking document intelligence systems on academic credentials remains severely "
        "bottlenecked by statutory privacy frameworks, structural diversity, and rigid evaluation "
        "metrics that unduly penalize benign formatting variations. To resolve these challenges, "
        "this paper presented an integrated end-to-end evaluation architecture comprising ADBG v1.0 "
        "and the AU DIC Framework v1.0. The framework incorporates seed-deterministic credential "
        "synthesis across certificates, marksheets, and student identity cards, a four-profile optical "
        "degradation matrix simulating real-world capture noise, a six-stage semantic canonical "
        "normalizer, and a structured nine-class diagnostic OCR error taxonomy. Live empirical evaluation "
        "on 450 multimodal specimens and 9,000 paired field observations confirmed that canonical "
        "normalization significantly elevates extraction fidelity, increasing Field F1 from 81.77% to "
        "92.92% and reducing Character Error Rate by 69.90% (from 8.14% to 2.45%). Rigorous statistical "
        "hypothesis testing validated these performance gains at p < 0.0001 (McNemar chi2 = 1002.00, "
        "Wilcoxon W = 0.0), establishing a robust, privacy-preserving standard for institutional credential intelligence."
    )

    fw_text = (
        "Promising avenues for future research encompass three primary dimensions to extend the benchmark "
        "diagnostic reach. First, future efforts will expand deterministic generation to multilingual and "
        "international academic credentials, supporting cross-jurisdictional transcript schemas, multi-alphabet "
        "typography, and non-Latin character sets. Second, the evaluation protocol will incorporate "
        "next-generation document foundation backbones alongside specialized vision-language models (such as "
        "LayoutLMv2 [41] and Qwen2.5-VL) to systematically assess model resilience under extreme optical capture "
        "distortions and complex physical degradation profiles. Third, future development will target multi-model "
        "ensemble architectures and automated confidence calibration pipelines to deliver fault-tolerant, "
        "production-ready document verification workflows for institutional administrative systems and automated credential auditing."
    )

    concl_words = len(concl_text.split())
    fw_words = len(fw_text.split())
    total_words = concl_words + fw_words
    print(f'Conclusion words: {concl_words}')
    print(f'Future Work words: {fw_words}')
    print(f'Combined words: {total_words}')
    assert total_words == 250, f'Expected 250 words, got {total_words}'

    # Update Conclusion body
    concl_p.runs[0].text = concl_text
    for r in concl_p.runs[1:]:
        r.text = ''
    concl_p.runs[0].font.name = 'Times New Roman'
    concl_p.runs[0].font.size = Pt(10.0)
    concl_p.runs[0].bold = None
    concl_p.runs[0].italic = None
    concl_p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    concl_p.paragraph_format.space_after = Pt(3.5)
    concl_p.paragraph_format.line_spacing = 1.0791666666666666

    # Update Future Work body
    fw_p.runs[0].text = fw_text
    for r in fw_p.runs[1:]:
        r.text = ''
    fw_p.runs[0].font.name = 'Times New Roman'
    fw_p.runs[0].font.size = Pt(10.0)
    fw_p.runs[0].bold = None
    fw_p.runs[0].italic = None
    fw_p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    fw_p.paragraph_format.space_after = Pt(3.5)
    fw_p.paragraph_format.line_spacing = 1.0791666666666666

    doc.save(docx_path)
    print(f'Successfully updated {docx_path} with 250-word Conclusion and Future Work!')

if __name__ == '__main__':
    update_conclusion_and_future_work('PaperV50.docx')
