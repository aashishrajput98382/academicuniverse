import docx

doc = docx.Document('PaperV50.docx')
p69 = doc.paragraphs[69]
p70 = doc.paragraphs[70]

print('P69:')
print('  text:', p69.text)
print('  style:', p69.style.name)
print('  alignment:', p69.alignment)
print('  runs:', [(r.text, r.font.name, r.font.size, r.bold) for r in p69.runs])

print('P70:')
print('  text:', p70.text)
print('  style:', p70.style.name)
print('  alignment:', p70.alignment)
print('  runs:', [(r.text, r.font.name, r.font.size, r.bold) for r in p70.runs])
