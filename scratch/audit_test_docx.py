import docx, re
from docx.oxml.text.paragraph import CT_P
from docx.oxml.table import CT_Tbl
from docx.text.paragraph import Paragraph
from docx.table import Table

doc = docx.Document('PaperV50_test.docx')

body_elements = []
tbl_count = 0
p_count = 0

for child in doc._element.body:
    if isinstance(child, CT_P):
        p = Paragraph(child, doc)
        body_elements.append(('P', p_count, p.text.strip()))
        p_count += 1
    elif isinstance(child, CT_Tbl):
        tbl = Table(child, doc)
        body_elements.append(('TBL', tbl_count, len(tbl.rows), len(tbl.columns)))
        tbl_count += 1

print(f'Total elements: {len(body_elements)}, Paragraphs: {p_count}, Tables: {tbl_count}')

table_references_audit = []

for idx, elem in enumerate(body_elements):
    if elem[0] == 'TBL':
        tbl_idx = elem[1]
        # Find preceding 3 non-empty paragraphs
        prec = []
        k = idx - 1
        while k >= 0 and len(prec) < 3:
            if body_elements[k][0] == 'P' and body_elements[k][2]:
                prec.append(body_elements[k])
            k -= 1
        prec.reverse()

        # Find succeeding 3 non-empty paragraphs
        succ = []
        k = idx + 1
        while k < len(body_elements) and len(succ) < 3:
            if body_elements[k][0] == 'P' and body_elements[k][2]:
                succ.append(body_elements[k])
            k += 1

        caption = prec[-1][2] if prec else ''
        above_text = prec[-2][2] if len(prec) >= 2 else ''
        below_text = succ[0][2] if succ else ''

        table_references_audit.append({
            'tbl_idx': tbl_idx,
            'caption': caption,
            'above_p': prec[-2] if len(prec) >= 2 else None,
            'below_p': succ[0] if succ else None,
        })

print('\n' + '='*80)
print('TABLE-BY-TABLE REFERENCE AUDIT')
print('='*80)

for item in table_references_audit:
    t_idx = item['tbl_idx']
    cap = item['caption']
    above = item['above_p'][2] if item['above_p'] else 'NONE'
    below = item['below_p'][2] if item['below_p'] else 'NONE'

    # Extract table number from caption if possible
    m = re.search(r'TABLE\s*(\d+)', cap, re.IGNORECASE)
    t_num = m.group(1) if m else str(t_idx + 1)

    has_ref_above = f'Table {t_num}' in above or f'Table{t_num}' in above
    has_ref_below = f'Table {t_num}' in below or f'Table{t_num}' in below
    has_ref_cap = f'Table {t_num}' in cap or f'Table{t_num}' in cap

    print(f'\n--- Table Index {t_idx} (Table {t_num}) ---')
    print(f'  Caption: {cap[:75]}')
    print(f'  Paragraph Above: {repr(above[:90])}')
    print(f'  Paragraph Below: {repr(below[:90])}')
    print(f'  Referenced in Above Paragraph: {has_ref_above}')
    print(f'  Referenced in Below Paragraph: {has_ref_below}')

# Verify table font sizes
table_sizes = set()
for tbl in doc.tables:
    for r in tbl.rows:
        for c in r.cells:
            for p in c.paragraphs:
                for run in p.runs:
                    table_sizes.add(run.font.size.pt if run.font.size else None)

print('\n' + '='*80)
print(f'Table font sizes across all tables: {table_sizes}')
assert table_sizes == {8.0}, f'Unexpected sizes: {table_sizes}'

# Verify 0 asterisks
p_stars = [i for i, p in enumerate(doc.paragraphs) if '*' in p.text]
tbl_stars = [t for t, tbl in enumerate(doc.tables) if any('*' in c.text for r in tbl.rows for c in r.cells)]
print(f'Asterisks count: {len(p_stars) + len(tbl_stars)}')
assert len(p_stars) == 0 and len(tbl_stars) == 0, 'Found asterisks!'

# Check Conclusion and Future Work word count
conclusion_text = ''
future_work_text = ''
c_active = False
f_active = False

for p in doc.paragraphs:
    txt = p.text.strip()
    if '6 CONCLUSION' in txt:
        c_active = True
        f_active = False
        continue
    elif '7 FUTURE WORK' in txt:
        c_active = False
        f_active = True
        continue
    elif 'DECLARATION' in txt:
        c_active = False
        f_active = False
        break
    
    if c_active and txt:
        conclusion_text += ' ' + txt
    elif f_active and txt:
        future_work_text += ' ' + txt

c_words = len(conclusion_text.split())
f_words = len(future_work_text.split())
print(f'Conclusion words: {c_words}, Future work words: {f_words}, Total: {c_words + f_words}')
assert c_words + f_words == 250, f'Conclusion + Future work count is {c_words + f_words}, expected 250!'

print('\nALL VERIFICATIONS PASSED PERFECTLY!')
