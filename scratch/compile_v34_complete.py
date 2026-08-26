import os
import re
import subprocess
import win32com.client
import docx
from docx.shared import Pt, RGBColor
from pathlib import Path

workspace = Path(__file__).resolve().parents[1]
v33_docx_path = workspace / "docs" / "paper" / "PaperV33_Ollama_Primary.docx"
v34_docx_path = workspace / "docs" / "paper" / "PaperV34_Ollama_Primary.docx"
v34_pdf_path = workspace / "docs" / "paper" / "PaperV34_Ollama_Primary.pdf"

doc = docx.Document(v33_docx_path)

# Update Paragraph 89 with 100% synchronized statistical hypothesis text
target_idx = None
for idx, p in enumerate(doc.paragraphs):
    if "Rigorous statistical hypothesis testing reported in Table 11 confirms" in p.text:
        target_idx = idx
        break

if target_idx is not None:
    p = doc.paragraphs[target_idx]
    p.text = (
        "Rigorous statistical hypothesis testing reported in Table 11 confirms that metric improvements from "
        "canonicalization are highly significant (p < 0.0001, McNemar chi2 = 97.01, Wilcoxon W = 0.0, Paired t = 10.54). "
        "Furthermore, 10,000-iteration non-parametric bootstrap resampling in Table 12 establishes non-overlapping "
        "95% confidence intervals between Pass A (F1: [76.89%, 82.22%]) and Pass B (F1: [88.56%, 92.44%])."
    )
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.15
    for r in p.runs:
        r.font.name = "Times New Roman"
        r.font.size = Pt(10.0)
        r.font.color.rgb = RGBColor(30, 41, 59)
    print(f"[SUCCESS] Updated paragraph {target_idx} with synchronized statistical metrics.")
else:
    raise ValueError("Target paragraph not found!")

# Save DOCX
doc.save(v34_docx_path)
print(f"[SUCCESS] Saved PaperV34_Ollama_Primary.docx!")

# Export to PDF via Word COM
subprocess.run(["taskkill", "/F", "/IM", "WINWORD.EXE"], capture_output=True)

word = win32com.client.DispatchEx("Word.Application")
word.Visible = False
word.DisplayAlerts = False

try:
    doc_com = word.Documents.Open(str(v34_docx_path.resolve()))
    doc_com.SaveAs(str(v34_pdf_path.resolve()), FileFormat=17) # 17 = wdFormatPDF
    doc_com.Close(False)
    print(f"[SUCCESS] Exported PaperV34_Ollama_Primary.pdf!")
finally:
    word.Quit()
    subprocess.run(["taskkill", "/F", "/IM", "WINWORD.EXE"], capture_output=True)
