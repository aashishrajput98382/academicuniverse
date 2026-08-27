import os
import re
import subprocess
import win32com.client
import docx
import pypdf
import fitz
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls
from pathlib import Path

workspace = Path(__file__).resolve().parents[1]
v44_docx_path = workspace / "docs" / "paper" / "PaperV44_Ollama_Primary.docx"
test_docx = workspace / "scratch" / "test_flow_20.docx"
test_pdf  = workspace / "scratch" / "test_flow_20.pdf"

doc = docx.Document(v44_docx_path)

# Image paths
fig1_png = workspace / "docs" / "paper" / "figure1_system_architecture_mermaid.png"
fig2_png = workspace / "docs" / "paper" / "figure2_data_flow_mermaid.png"
if not fig2_png.exists():
    fig2_png = workspace / "docs" / "paper" / "methodology_workflow_600dpi.png"

extracted_dir = workspace / "docs" / "paper" / "extracted_figures"
fig3_path = extracted_dir / "image3.png"
fig4_path = extracted_dir / "image4.png"
fig5_path = extracted_dir / "image5.png"
fig6_path = extracted_dir / "image6.png"
fig7_path = extracted_dir / "image7.png"
fig8_path = workspace / "results" / "confusion_matrices" / "dt_composite.png"
fig9_path = workspace / "results" / "confusion_matrices" / "rf_composite.png"

# 1. Direct replacements for text & tables
replacements = [
    ("χ² = 97.01, p < 0.0001", "χ² = 1,002.00, p < 10⁻²⁰⁰"),
    ("χ² = 97.01", "χ² = 1,002.00"),
    ("97.01", "1002.00"),
    ("79.56%", "81.77%"),
    ("90.56%", "92.92%"),
    ("9.69% to 3.64%", "8.14% to 2.45%"),
    ("9.69%", "8.14%"),
    ("3.64%", "2.45%"),
    ("62.44% relative CER reduction", "69.90% relative CER reduction"),
    ("62.44%", "69.90%"),
    ("+11.00% net gain", "+11.16% net gain"),
    ("+11.00%", "+11.16%"),
    ("11.00%", "11.16%"),
    ("+13.83% relative", "+13.64% relative"),
    ("+13.83%", "+13.64%"),
    ("-62.44%", "-69.90%"),
    ("-6.05%", "-5.69%"),
    ("6.05%", "5.69%"),
    ("97.80% student name recognition, 93.30% university recognition", "96.00% student name recognition, 79.10% university recognition"),
    ("97.80% student name", "96.00% student name"),
    ("93.30% university", "79.10% university"),
    ("across 900 paired field observations", "across 9,000 paired field observations"),
    ("across 900 atomic field observations", "across 9,000 atomic field observations"),
    ("(900 paired field observations)", "(9,000 paired field observations)"),
    ("900 evaluated ground-truth field observations", "9,000 evaluated ground-truth field observations"),
    ("900 field observations", "9,000 field observations"),
    ("900 paired extractions", "9,000 paired extractions"),
    ("900 candidate field observations", "9,000 candidate field observations"),
    ("900 observations", "9,000 observations"),
    ("(45 Specimens / 900 Fields)", "(450 Specimens / 9,000 Fields)"),
    ("(N = 900,", "(N = 9,000,"),
    ("M = 900", "M = 9,000"),
    ("900 (100%)", "9,000 (100%)"),
    ("99 false-negative field mismatches", "1,004 false-negative field mismatches"),
    ("99 false-negative formatting mismatches", "1,004 false-negative formatting mismatches"),
    ("resolves 99 false-negative", "resolves 1,004 false-negative"),
    ("resolved 99 false-negative", "resolved 1,004 false-negative"),
    ("All 99 `FORMAT_ERROR`", "All 1,004 `FORMAT_ERROR`"),
    ("All 99 FORMAT_ERROR", "All 1,004 FORMAT_ERROR"),
    ("raising exact match from 79.56% to 90.56%", "raising exact match from 81.77% to 92.92%"),
    ("815 `EXACT_MATCH`", "8,363 `EXACT_MATCH`"),
    ("815 EXACT_MATCH", "8,363 EXACT_MATCH"),
    ("716 (79.56%)", "7,359 (81.77%)"),
    ("815 (90.56%)", "8,363 (92.92%)"),
    ("99 (11.00%)", "1,004 (11.16%)"),
    ("85 (9.44%)", "637 (7.08%)"),
    ("run_1785959173886", "run_450_live_gpu_1787773778444"),
    ("2917.90s / 48.63 mins", "35746.20s / 595.77 mins"),
    ("2917.90s", "35746.20s"),
    ("48.63 mins", "595.77 mins"),
    ("0bc63b0", "18745ce")
]

# Run-level replacements in paragraphs
for p in doc.paragraphs:
    for r in p.runs:
        for old_s, new_s in replacements:
            if old_s in r.text:
                r.text = r.text.replace(old_s, new_s)

# Run-level replacements in tables
for tbl in doc.tables:
    for row in tbl.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                for r in p.runs:
                    for old_s, new_s in replacements:
                        if old_s in r.text:
                            r.text = r.text.replace(old_s, new_s)

# 2. Perfect image widths for natural single-page fit
for i, p in enumerate(doc.paragraphs):
    if p.text.strip().startswith("Fig. 1.") and "System Architecture" in p.text:
        img_p = doc.paragraphs[i - 1]
        img_p.clear()
        img_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        img_p.paragraph_format.space_before = Pt(1)
        img_p.paragraph_format.space_after = Pt(1)
        img_p.paragraph_format.keep_with_next = True
        r = img_p.add_run()
        r.add_picture(str(fig1_png.resolve()), width=Inches(2.7))
        p.paragraph_format.keep_with_next = True
        p.paragraph_format.keep_lines = True
        p.paragraph_format.space_before = Pt(1)
        p.paragraph_format.space_after = Pt(2)

    if p.text.strip().startswith("Fig. 2.") and "Data Flow" in p.text:
        img_p = doc.paragraphs[i - 1]
        img_p.clear()
        img_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        img_p.paragraph_format.space_before = Pt(1)
        img_p.paragraph_format.space_after = Pt(1)
        img_p.paragraph_format.keep_with_next = True
        r = img_p.add_run()
        r.add_picture(str(fig2_png.resolve()), width=Inches(2.5))
        p.paragraph_format.keep_with_next = True
        p.paragraph_format.keep_lines = True
        p.paragraph_format.space_before = Pt(1)
        p.paragraph_format.space_after = Pt(2)

# Size confusion matrix composite images
for i, p in enumerate(doc.paragraphs):
    if "Confusion matrices for Decision Tree" in p.text and fig8_path.exists():
        img_p = doc.paragraphs[i - 1]
        img_p.clear()
        img_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        img_p.paragraph_format.keep_with_next = True
        r = img_p.add_run()
        r.add_picture(str(fig8_path.resolve()), width=Inches(2.8))

    if "Confusion matrices for Random Forest" in p.text and fig9_path.exists():
        img_p = doc.paragraphs[i - 1]
        img_p.clear()
        img_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        img_p.paragraph_format.keep_with_next = True
        r = img_p.add_run()
        r.add_picture(str(fig9_path.resolve()), width=Inches(2.8))

# General paragraph spacing adjustments
for p in doc.paragraphs:
    if p.text.strip().startswith("TABLE "):
        p.paragraph_format.keep_with_next = True
        p.paragraph_format.keep_lines = True
        p.paragraph_format.space_before = Pt(2.5)
        p.paragraph_format.space_after = Pt(1)
    elif p.text.strip().startswith("Fig. ") or p.text.strip().startswith("Figure "):
        p.paragraph_format.keep_lines = True
        p.paragraph_format.space_before = Pt(1)
        p.paragraph_format.space_after = Pt(2)

# Ensure cantSplit on all table rows
for tbl in doc.tables:
    for row in tbl.rows:
        trPr = row._tr.get_or_add_trPr()
        trPr.append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>'))

# References formatting
ref_started = False
for p in doc.paragraphs:
    if "REFERENCES" in p.text:
        ref_started = True
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(1)
        continue
    if ref_started and re.match(r"^\[\d+\]", p.text.strip()):
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0.25)
        p.paragraph_format.line_spacing = 1.0

doc.save(test_docx)
print("[SUCCESS] Saved test_flow_20.docx")

# Export to PDF via Word COM
subprocess.run(["taskkill", "/F", "/IM", "WINWORD.EXE"], capture_output=True)
word = win32com.client.DispatchEx("Word.Application")
word.Visible = False
word.DisplayAlerts = False
try:
    doc_com = word.Documents.Open(str(test_docx.resolve()))
    doc_com.SaveAs(str(test_pdf.resolve()), FileFormat=17)
    doc_com.Close(False)
    print("[SUCCESS] Exported test_flow_20.pdf")
finally:
    word.Quit()
    subprocess.run(["taskkill", "/F", "/IM", "WINWORD.EXE"], capture_output=True)

# Verify page count
reader = pypdf.PdfReader(str(test_pdf))
pdf_pages = len(reader.pages)
print(f"\n[INFO] Page Count: {pdf_pages} pages")
