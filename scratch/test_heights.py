import docx
from docx.shared import Inches
import pymupdf
import win32com.client, os
import shutil

src_backup = r'c:\github\academicuniverse\ashmitxnaik\Deepfake_Video_Authentication_Updated_Research_Paper_With_Flowchart (1)_backup.docx'
img_path = r'c:\github\academicuniverse\ashmitxnaik\deepfake_architecture_flowchart.png'

word = win32com.client.Dispatch('Word.Application')
word.Visible = False

for test_h in [1.5, 1.8, 2.0, 2.2, 2.5]:
    test_docx = rf'c:\github\academicuniverse\scratch\test_h_{test_h}.docx'
    test_pdf = rf'c:\github\academicuniverse\scratch\test_h_{test_h}.pdf'
    shutil.copyfile(src_backup, test_docx)
    
    doc = docx.Document(test_docx)
    # find target
    t_start = None
    for i, p in enumerate(doc.paragraphs):
        if 'Input Video' in p.text and 'Facial Sequence' in p.text:
            t_start = i
            break
            
    to_rem = []
    for j in range(t_start + 1, t_start + 15):
        pj = doc.paragraphs[j]
        to_rem.append(pj)
        if 'Authentication Dashboard' in pj.text:
            break
            
    p_img = doc.paragraphs[t_start]
    p_img.text = ""
    p_img.alignment = docx.enum.text.WD_ALIGN_PARAGRAPH.CENTER
    r = p_img.add_run()
    # add with specific height
    r.add_picture(img_path, height=Inches(test_h))
    
    # Caption
    p_cap = to_rem[0]
    p_cap.text = "Fig. 1. Architecture of the proposed framework."
    p_cap.alignment = docx.enum.text.WD_ALIGN_PARAGRAPH.CENTER
    for run in p_cap.runs:
        run.font.size = 101600 # 8 pt
        run.italic = True
        
    for pj in to_rem[1:]:
        pe = pj._p
        pe.getparent().remove(pe)
        
    doc.save(test_docx)
    
    d = word.Documents.Open(os.path.abspath(test_docx))
    d.SaveAs(os.path.abspath(test_pdf), FileFormat=17)
    d.Close()
    
    pdf = pymupdf.open(test_pdf)
    print(f"Height {test_h} in -> Total pages: {len(pdf)}")
    # Check if 'TEMPORAL FEATURE LEARNING' is on page 3
    p3_text = pdf[2].get_text() # page 3
    has_ix = 'TEMPORAL FEATURE LEARNING' in p3_text
    print(f"  Page 3 has TEMPORAL FEATURE LEARNING: {has_ix}")

word.Quit()
