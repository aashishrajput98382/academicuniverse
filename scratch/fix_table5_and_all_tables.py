import docx
import re
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

# 1. Update Paper_V22.md
with open('docs/paper/Paper_V22.md', 'r', encoding='utf-8') as f:
    md_text = f.read()

# Update Table 5 in md_text if needed
md_text = md_text.replace('24,480 observations', '900 field observations across 45 physical specimens')
md_text = md_text.replace('360-specimen synthetic evaluation dataset', 'Canonical 45-specimen multi-modal evaluation dataset')
md_text = md_text.replace('360 paired image and JSON instances', '45 paired image and JSON instances')
md_text = md_text.replace('24,480 paired field observations', '900 paired field observations')

with open('docs/paper/Paper_V22.md', 'w', encoding='utf-8') as f:
    f.write(md_text)

print('[SUCCESS] Updated Paper_V22.md')

# 2. Update scratch/generate_paperv22_pipeline.py with Table 5 and Table 1 fixes
with open('scratch/generate_paperv22_pipeline.py', 'r', encoding='utf-8') as f:
    pipe_code = f.read()

# Replace Table 5 / doc.tables[3] values
table_5_fix = '''
# Update Table 1 (Related Work Strategy cell)
if len(doc.tables) > 0 and len(doc.tables[0].rows) > 10:
    for cell in doc.tables[0].rows[10].cells:
        if "24,480" in cell.text:
            cell.text = cell.text.replace("24,480", "900")

# Update Table 4 / 5 (Canonical Configuration Parameters)
if len(doc.tables) > 3:
    t5 = doc.tables[3]
    for r_idx, row in enumerate(t5.rows):
        param_name = row.cells[0].text.strip()
        if "Total Evaluated Specimens" in param_name:
            row.cells[1].text = "45 Physical Specimens (5 Vector PDFs + 20 Lossless PNGs + 20 Compressed JPEGs)"
        elif "Total Paired Observations" in param_name:
            row.cells[1].text = "900 Paired Field Observations (20 Atomic Fields / Specimen)"
        elif "Canonical Execution Run ID" in param_name:
            row.cells[1].text = "run_1785959173886 (Duration: 2917.90s / 48.63 mins)"
        elif "Git Repository Commit" in param_name:
            row.cells[1].text = "Commit 0cb27be (https://github.com/aashishrajput9838/academicuniverse.git)"
        elif "Benchmark Suite Version" in param_name:
            row.cells[1].text = "AU DIC Benchmark v1.0 (45 Physical Specimens Suite)"
'''

if "t5 = doc.tables[3]" not in pipe_code:
    pipe_code = pipe_code.replace('doc.save(v22_docx_path)', table_5_fix + '\ndoc.save(v22_docx_path)')

with open('scratch/generate_paperv22_pipeline.py', 'w', encoding='utf-8') as f:
    f.write(pipe_code)

print('[SUCCESS] Injected Table 5 and Table 1 fixes into scratch/generate_paperv22_pipeline.py!')
