import docx
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
import shutil

src = r'c:\github\academicuniverse\ashmitxnaik\Deepfake_Video_Authentication_Updated_Research_Paper_With_Flowchart (1).docx'
backup = r'c:\github\academicuniverse\ashmitxnaik\Deepfake_Video_Authentication_Updated_Research_Paper_With_Flowchart (1)_backup.docx'
dst = r'c:\github\academicuniverse\scratch\test_applied_paper.docx'

# Create backup
shutil.copyfile(src, backup)
shutil.copyfile(src, dst)
print("Created backup at:", backup)

doc = docx.Document(dst)

# Identify target paragraphs
target_start = None
for i, p in enumerate(doc.paragraphs):
    if 'Input Video' in p.text and 'Facial Sequence Construction' in p.text:
        target_start = i
        break

print(f"Target start: {target_start}")
to_remove = []
for j in range(target_start + 1, target_start + 15):
    pj = doc.paragraphs[j]
    to_remove.append(pj)
    if 'Authentication Dashboard' in pj.text:
        break

print(f"Paragraphs to remove count: {len(to_remove)}")
for pj in to_remove:
    print(f"  Removing: {repr(pj.text)}")

# Replace target_start paragraph with image
p_img = doc.paragraphs[target_start]
p_img.text = ""
p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
run_img = p_img.add_run()
img_path = r'c:\github\academicuniverse\ashmitxnaik\deepfake_architecture_flowchart.png'
# Column width is 5602 twips = 3.89 inches. Width of 3.48 inches fits cleanly inside column with margin.
run_img.add_picture(img_path, width=Inches(3.48))

# Caption paragraph right after the image
# Let's see: we can use the first removed paragraph or insert one
p_cap = to_remove[0] # use the first one
p_cap.text = "Fig. 1. Overall architecture of the proposed explainable temporal deepfake authentication framework."
p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
if p_cap.runs:
    for r in p_cap.runs:
        r.font.size = 114300 # 9 pt
        r.italic = True

# Remove remaining paragraphs
for pj in to_remove[1:]:
    p_elem = pj._p
    p_elem.getparent().remove(p_elem)

doc.save(dst)
print("Saved modified test paper to:", dst)
