import os
import re
import subprocess
import win32com.client
import docx
from docx.shared import Pt, RGBColor
from pathlib import Path

workspace = Path(__file__).resolve().parents[1]
v34_docx_path = workspace / "docs" / "paper" / "PaperV34_Ollama_Primary.docx"
v35_docx_path = workspace / "docs" / "paper" / "PaperV35_Ollama_Primary.docx"
v35_pdf_path = workspace / "docs" / "paper" / "PaperV35_Ollama_Primary.pdf"

doc = docx.Document(v34_docx_path)

# 1. Update Table 8 Heading
for idx, p in enumerate(doc.paragraphs):
    if "TABLE 8: STATE-OF-THE-ART" in p.text.upper() or "TABLE 8:" in p.text.upper():
        if "BASELINE" in p.text.upper() or "BENCHMARK" in p.text.upper():
            p.text = "TABLE 8: CONTEXTUAL COMPARISON WITH REPORTED DOCUMENT INTELLIGENCE BASELINES ACROSS EVALUATION PARADIGMS"
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(2)
            for r in p.runs:
                r.font.name = "Times New Roman"
                r.font.size = Pt(9.0)
                r.font.bold = True
                r.font.color.rgb = RGBColor(27, 54, 93)
            print(f"[SUCCESS] Updated Table 8 Heading at paragraph {idx}")
            break

# 2. Update Table 8 (Table 6 in docx) Row 5 last cell
for t_idx, table in enumerate(doc.tables):
    for r_idx, row in enumerate(table.rows):
        row_txt = " | ".join(c.text.strip() for c in row.cells)
        if "Proposed AU DIC System (Ours)" in row_txt and "Live Verified State-of-the-Art" in row_txt:
            row.cells[-1].text = "Best performance observed within evaluated AU DIC configuration (+11.00% net gain)"
            for p in row.cells[-1].paragraphs:
                for r in p.runs:
                    r.font.name = "Times New Roman"
                    r.font.size = Pt(8.0)
                    r.font.bold = True
            print(f"[SUCCESS] Updated Table {t_idx} Row {r_idx} last cell")

# Save DOCX
doc.save(v35_docx_path)
print(f"[SUCCESS] Saved PaperV35_Ollama_Primary.docx!")

# Export to PDF via Word COM
subprocess.run(["taskkill", "/F", "/IM", "WINWORD.EXE"], capture_output=True)

word = win32com.client.DispatchEx("Word.Application")
word.Visible = False
word.DisplayAlerts = False

try:
    doc_com = word.Documents.Open(str(v35_docx_path.resolve()))
    doc_com.SaveAs(str(v35_pdf_path.resolve()), FileFormat=17) # 17 = wdFormatPDF
    doc_com.Close(False)
    print(f"[SUCCESS] Exported PaperV35_Ollama_Primary.pdf!")
finally:
    word.Quit()
    subprocess.run(["taskkill", "/F", "/IM", "WINWORD.EXE"], capture_output=True)
