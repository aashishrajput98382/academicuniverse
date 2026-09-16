import docx
from docx.shared import Inches
import pymupdf
import win32com.client, os
import shutil

src_backup = r'c:\github\academicuniverse\ashmitxnaik\Deepfake_Video_Authentication_Updated_Research_Paper_With_Flowchart (1)_backup.docx'
test_docx = r'c:\github\academicuniverse\scratch\test_horizontal.docx'
test_pdf = r'c:\github\academicuniverse\scratch\test_horizontal.pdf'
img_path = r'c:\github\academicuniverse\ashmitxnaik\deepfake_architecture_flowchart_horizontal.png'

shutil.copyfile(src_backup, test_docx)

doc = docx.Document(test_docx)
t_start = None
for i, p in enumerate(doc.paragraphs):
    if 'Input Video' in p.text and 'Facial Sequence' in p.text:
        t_start = i
        break

to_rem = []
for j in range(t_start + 1, t_start + 15):
    pj = doc.paragraphs[j]
    to_rem.append(pj)
    if 'Authentication Dashboard' in pj.text:
        break

p_img = doc.paragraphs[t_start]
p_img.text = ""
p_img.alignment = docx.enum.text.WD_ALIGN_PARAGRAPH.CENTER
r = p_img.add_run()
r.add_picture(img_path, width=Inches(3.65))

# Caption
p_cap = to_rem[0]
p_cap.text = "Fig. 1. Overall architecture of the proposed explainable temporal deepfake authentication framework."
p_cap.alignment = docx.enum.text.WD_ALIGN_PARAGRAPH.CENTER
for run in p_cap.runs:
    run.font.size = 101600 # 8 pt
    run.italic = True

for pj in to_rem[1:]:
    pe = pj._p
    pe.getparent().remove(pe)

doc.save(test_docx)

word = win32com.client.Dispatch('Word.Application')
word.Visible = False
d = word.Documents.Open(os.path.abspath(test_docx))
d.SaveAs(os.path.abspath(test_pdf), FileFormat=17)
d.Close()
word.Quit()

pdf = pymupdf.open(test_pdf)
print(f"Total pages in horizontal test: {len(pdf)}")
for pno in range(len(pdf)):
    pix = pdf[pno].get_pixmap(dpi=150)
    pix.save(f"scratch/hz_page_{pno+1}.png")
    print(f"Rendered scratch/hz_page_{pno+1}.png")
