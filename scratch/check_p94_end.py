import docx, re

doc = docx.Document('PaperV50.docx')
for i in range(94, len(doc.paragraphs)):
    txt = doc.paragraphs[i].text.strip()
    if re.search(r'table', txt, re.I):
        print(f'P{i}: {txt}')
