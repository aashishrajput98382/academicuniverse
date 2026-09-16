import docx
from docx.shared import Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
import win32com.client, os, shutil

# Source files
src_backup = r'c:\github\academicuniverse\ashmitxnaik\Deepfake_Video_Authentication_Updated_Research_Paper_With_Flowchart (1)_backup.docx'
target_docx = r'c:\github\academicuniverse\ashmitxnaik\Deepfake_Video_Authentication_Updated_Research_Paper_With_Flowchart (1).docx'
target_pdf = r'c:\github\academicuniverse\ashmitxnaik\Deepfake_Video_Authentication_Updated_Research_Paper_With_Flowchart (1).pdf'
img_horizontal = r'c:\github\academicuniverse\ashmitxnaik\deepfake_architecture_flowchart_horizontal.png'
img_primary = r'c:\github\academicuniverse\ashmitxnaik\deepfake_architecture_flowchart.png'

# Also update the primary image file with the horizontal layout
shutil.copyfile(img_horizontal, img_primary)

# Start clean from backup to avoid any accumulated edits
shutil.copyfile(src_backup, target_docx)

doc = docx.Document(target_docx)

# Locate target block
t_start = None
for i, p in enumerate(doc.paragraphs):
    if 'Input Video' in p.text and 'Facial Sequence' in p.text:
        t_start = i
        break

if t_start is None:
    raise ValueError("Target paragraph not found!")

to_rem = []
for j in range(t_start + 1, t_start + 15):
    pj = doc.paragraphs[j]
    to_rem.append(pj)
    if 'Authentication Dashboard' in pj.text:
        break

# Replace target start with image
p_img = doc.paragraphs[t_start]
p_img.text = ""
p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p_img.add_run()
# 3.65 inches wide fits cleanly inside column (3.89 inches) without spilling vertically
r.add_picture(img_primary, width=Inches(3.65))

# Caption
p_cap = to_rem[0]
p_cap.text = "Fig. 1. Overall architecture of the proposed explainable temporal deepfake authentication framework."
p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
for run in p_cap.runs:
    run.font.size = 101600 # 8 pt
    run.italic = True

# Remove remaining flowchart text paragraphs
for pj in to_rem[1:]:
    pe = pj._p
    pe.getparent().remove(pe)

doc.save(target_docx)
print("Updated target docx successfully.")

# Recompile PDF
word = win32com.client.Dispatch('Word.Application')
word.Visible = False
d = word.Documents.Open(os.path.abspath(target_docx))
d.SaveAs(os.path.abspath(target_pdf), FileFormat=17)
d.Close()
word.Quit()
print("Compiled target PDF successfully at:", target_pdf)
