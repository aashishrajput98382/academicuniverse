import docx
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

def apply_declaration_and_code_repository(docx_path):
    doc = docx.Document(docx_path)

    # Locate ACKNOWLEDGMENT and REFERENCES
    p_ack_head = None
    p_ack_body = None
    p_ref = None

    for i, p in enumerate(doc.paragraphs):
        t = p.text.strip().upper()
        if t == 'ACKNOWLEDGMENT':
            p_ack_head = p
            p_ack_body = doc.paragraphs[i+1]
        elif t == 'REFERENCES':
            p_ref = p
            break

    assert p_ack_head is not None, 'ACKNOWLEDGMENT not found!'
    assert p_ref is not None, 'REFERENCES not found!'

    # 1. Replace ACKNOWLEDGMENT heading with DECLARATION
    p_ack_head.runs[0].text = 'DECLARATION'
    for r in p_ack_head.runs[1:]:
        r.text = ''
    p_ack_head.runs[0].font.name = 'Times New Roman'
    p_ack_head.runs[0].font.size = Pt(10.0)
    p_ack_head.runs[0].bold = True
    p_ack_head.runs[0].italic = None
    p_ack_head.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p_ack_head.paragraph_format.space_after = Pt(3.5)
    p_ack_head.paragraph_format.line_spacing = 1.0791666666666666

    # 2. Replace Acknowledgment body with Declaration body
    decl_text = (
        "Conflict of Interest: The authors declare that they have no conflicts of interest "
        "or competing financial interests to declare. Ethical Approval: Not applicable; this "
        "study was conducted exclusively using synthetic benchmark specimens generated without "
        "human participants or personally identifiable educational records. Funding: The authors "
        "declare that no external grants, funds, or other financial support were received for the "
        "research, authorship, or publication of this work."
    )
    p_ack_body.runs[0].text = decl_text
    for r in p_ack_body.runs[1:]:
        r.text = ''
    p_ack_body.runs[0].font.name = 'Times New Roman'
    p_ack_body.runs[0].font.size = Pt(10.0)
    p_ack_body.runs[0].bold = None
    p_ack_body.runs[0].italic = None
    p_ack_body.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p_ack_body.paragraph_format.space_after = Pt(3.5)
    p_ack_body.paragraph_format.line_spacing = 1.0791666666666666

    # 3. Insert CODE REPOSITORY heading and body before REFERENCES
    p_code_head = p_ref.insert_paragraph_before('CODE REPOSITORY')
    p_code_head.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p_code_head.paragraph_format.space_after = Pt(3.5)
    p_code_head.paragraph_format.line_spacing = 1.0791666666666666
    r_ch = p_code_head.runs[0]
    r_ch.font.name = 'Times New Roman'
    r_ch.font.size = Pt(10.0)
    r_ch.bold = True
    r_ch.italic = None

    code_text = (
        "The complete benchmark evaluation framework, synthetic academic credential generator "
        "(ADBG v1.0), and reproducible evaluation artifacts are publicly available under "
        "open-access version control in the official project repository: "
        "https://github.com/aashishrajput9838/academicuniverse.git"
    )
    p_code_body = p_ref.insert_paragraph_before(code_text)
    p_code_body.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p_code_body.paragraph_format.space_after = Pt(3.5)
    p_code_body.paragraph_format.line_spacing = 1.0791666666666666
    r_cb = p_code_body.runs[0]
    r_cb.font.name = 'Times New Roman'
    r_cb.font.size = Pt(10.0)
    r_cb.bold = None
    r_cb.italic = None

    doc.save(docx_path)
    print(f'Successfully updated {docx_path}!')

if __name__ == '__main__':
    apply_declaration_and_code_repository('PaperV50.docx')
