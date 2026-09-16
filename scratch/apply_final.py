import docx
from docx.shared import Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
import os

target_file = r'c:\github\academicuniverse\ashmitxnaik\Deepfake_Video_Authentication_Updated_Research_Paper_With_Flowchart (1).docx'
img_path = r'c:\github\academicuniverse\ashmitxnaik\deepfake_architecture_flowchart.png'

doc = docx.Document(target_file)

# Find target start paragraph
target_start = None
for i, p in enumerate(doc.paragraphs):
    if 'Input Video' in p.text and 'Facial Sequence Construction' in p.text:
        target_start = i
        break

if target_start is None:
    raise ValueError("Target paragraph not found in document!")

print(f"Target paragraph found at index: {target_start}")

# Gather paragraphs to remove
to_remove = []
for j in range(target_start + 1, target_start + 15):
    pj = doc.paragraphs[j]
    to_remove.append(pj)
    if 'Authentication Dashboard' in pj.text:
        break

print(f"Removing {len(to_remove)} flowchart text paragraphs.")

# Replace target_start paragraph with image
p_img = doc.paragraphs[target_start]
p_img.text = ""
p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
run_img = p_img.add_run()
run_img.add_picture(img_path, width=Inches(3.48))

# Caption paragraph right after the image
p_cap = to_remove[0]
p_cap.text = "Fig. 1. Overall architecture of the proposed explainable temporal deepfake authentication framework."
p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
if p_cap.runs:
    for r in p_cap.runs:
        r.font.size = 114300 # 9 pt
        r.italic = True

# Remove remaining text flowchart paragraphs
for pj in to_remove[1:]:
    p_elem = pj._p
    p_elem.getparent().remove(p_elem)

doc.save(target_file)
print("Successfully updated target document:", target_file)
