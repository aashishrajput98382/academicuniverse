import os
import re
import subprocess
import win32com.client
import docx
from docx.shared import Pt
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls
from pathlib import Path

workspace = Path(__file__).resolve().parents[1]
v23_docx_path = workspace / "docs" / "paper" / "PaperV23_Ollama_Primary.docx"
v23_pdf_path = workspace / "docs" / "paper" / "PaperV23_Ollama_Primary.pdf"
fig_dir = workspace / "docs" / "paper" / "extracted_figures"

# 1. Run pipeline
subprocess.run(["python", "scratch/generate_paperv23_pipeline.py"], check=True)

# 2. Replace figure image blobs with 300 DPI plots and crystal-clear flowcharts
doc = docx.Document(v23_docx_path)
replaced_count = 0
for r_id, rel in doc.part.rels.items():
    ref = rel.target_ref
    for img_name in ["image1.png", "image2.png", "image3.png", "image4.png", "image5.png", "image6.png", "image7.png", "image8.png", "image9.png"]:
        if img_name in ref:
            img_file = fig_dir / img_name
            if img_file.exists():
                with open(img_file, "rb") as f:
                    rel.target_part._blob = f.read()
                replaced_count += 1
                print(f"[REPLACED] {r_id} ({ref}) -> {img_name}")
print(f"[INFO] Total replaced image blobs: {replaced_count}")

# 3. Compact tables and drawings for <= 20 pages
for tbl in doc.tables:
    for row in tbl.rows:
        for cell in row.cells:
            tcPr = cell._tc.get_or_add_tcPr()
            tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="25" w:type="dxa"/><w:bottom w:w="25" w:type="dxa"/><w:left w:w="35" w:type="dxa"/><w:right w:w="35" w:type="dxa"/></w:tcMar>')
            tcPr.append(tcMar)
            for p in cell.paragraphs:
                p.paragraph_format.space_before = Pt(1)
                p.paragraph_format.space_after = Pt(1)
                p.paragraph_format.line_spacing = 1.0
                for r in p.runs:
                    r.font.name = "Times New Roman"
                    r.font.size = Pt(8.0)

for p in doc.paragraphs:
    txt = p.text.strip()
    if txt.startswith("TABLE ") or txt.startswith("Table "):
        p.paragraph_format.space_before = Pt(3)
        p.paragraph_format.space_after = Pt(1)
    elif txt.startswith("Fig. ") or txt.startswith("FIG. "):
        p.paragraph_format.space_before = Pt(1)
        p.paragraph_format.space_after = Pt(3)
    elif re.match(r'^\d+\.\s+', txt):
        p.paragraph_format.space_before = Pt(6)
        p.paragraph_format.space_after = Pt(2)
    else:
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.line_spacing = 1.05

    for r in p.runs:
        drawings = r._r.xpath('.//w:drawing')
        if drawings:
            for d in drawings:
                exts = d.xpath('.//wp:extent')
                for ext in exts:
                    cx = int(ext.get('cx', 0))
                    cy = int(ext.get('cy', 0))
                    if cy > 3.6 * 914400:
                        new_cy = int(3.6 * 914400)
                        new_cx = int(cx * (new_cy / cy))
                        ext.set('cx', str(new_cx))
                        ext.set('cy', str(new_cy))
                a_exts = d.xpath('.//a:ext')
                for a_ext in a_exts:
                    cx = int(a_ext.get('cx', 0))
                    cy = int(a_ext.get('cy', 0))
                    if cy > 3.6 * 914400:
                        new_cy = int(3.6 * 914400)
                        new_cx = int(cx * (new_cy / cy))
                        a_ext.set('cx', str(new_cx))
                        a_ext.set('cy', str(new_cy))

doc.save(v23_docx_path)
print(f"[SUCCESS] Saved compacted PaperV23_Ollama_Primary.docx!")

# 4. Export PDF
word = None
try:
    word = win32com.client.Dispatch("Word.Application")
    word.Visible = False
    doc_word = word.Documents.Open(os.path.abspath(str(v23_docx_path)))
    doc_word.SaveAs(os.path.abspath(str(v23_pdf_path)), FileFormat=17)
    page_count = doc_word.ComputeStatistics(2)
    doc_word.Close()
    print(f"[SUCCESS] Exported high-quality PDF: {v23_pdf_path.name} ({page_count} Pages!)")
except Exception as e:
    print(f"Word COM PDF Export Error: {e}")
finally:
    if word:
        try: word.Quit()
        except: pass
