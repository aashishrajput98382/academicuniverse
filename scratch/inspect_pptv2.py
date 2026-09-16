import pptx
from pptx.enum.shapes import MSO_SHAPE_TYPE

prs = pptx.Presentation('pptv2.pptx')
print(f"Presentation has {len(prs.slides)} slides. Dimensions: {prs.slide_width.pt} x {prs.slide_height.pt}")

for idx, slide in enumerate(prs.slides, 1):
    print(f"\n==================== SLIDE {idx} ====================")
    for s_idx, shape in enumerate(slide.shapes, 1):
        line = f"Shape {s_idx}: name='{shape.name}', type={shape.shape_type}, left={shape.left.pt:.1f}, top={shape.top.pt:.1f}, w={shape.width.pt:.1f}, h={shape.height.pt:.1f}"
        
        # Check fill
        try:
            fill = shape.fill
            if fill.type:
                line += f", fill={fill.type}"
                if fill.type == 1: # solid
                    line += f" rgb={fill.fore_color.rgb}"
        except Exception:
            pass
            
        print(line)
        if shape.has_text_frame:
            for p in shape.text_frame.paragraphs:
                p_text = p.text.strip()
                if p_text:
                    runs_info = []
                    for r in p.runs:
                        font_color = ""
                        try:
                            if r.font.color and r.font.color.rgb:
                                font_color = f"#{r.font.color.rgb}"
                        except Exception:
                            pass
                        runs_info.append(f"'{r.text}' (sz={r.font.size.pt if r.font.size else 'def'}, col={font_color}, b={r.font.bold})")
                    print(f"    P [level={p.level}]: {p_text}")
                    # print(f"       runs: {' + '.join(runs_info)}")
