
import pptx
from pptx.util import Pt, Inches
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE
import os
import win32com.client

# Design Tokens (White, Cream, Black)
C_CANVAS_WHITE = RGBColor(0xFF, 0xFF, 0xFF)
C_CARD_CREAM   = RGBColor(0xF7, 0xF3, 0xEB)
C_CARD_BORDER  = RGBColor(0xDF, 0xD7, 0xCB)
C_DIVIDER      = RGBColor(0xE5, 0xDD, 0xD0)
C_JET_BLACK    = RGBColor(0x11, 0x11, 0x11)
C_CHARCOAL     = RGBColor(0x24, 0x22, 0x20)
C_MUTED        = RGBColor(0x63, 0x5D, 0x56)
C_WHITE        = RGBColor(0xFF, 0xFF, 0xFF)
C_ACCENT_BG    = RGBColor(0x18, 0x18, 0x18)
C_CREAM_TEXT   = RGBColor(0xF7, 0xF3, 0xEB)
C_LIGHT_MUTED  = RGBColor(0x8A, 0x82, 0x78)
FONT_NAME      = "Segoe UI"

prs = pptx.Presentation("pptv2.pptx")
slide = prs.slides[0]

# 1. Slide 1 Canvas & Top Bar
slide.shapes[0].fill.solid()
slide.shapes[0].fill.fore_color.rgb = C_CANVAS_WHITE
slide.shapes[1].fill.solid()
slide.shapes[1].fill.fore_color.rgb = C_JET_BLACK

# 2. Hero Container (Rounded Rectangle 3)
hero = slide.shapes[2]
hero.left = Pt(40)
hero.top = Pt(25)
hero.width = Pt(880)
hero.height = Pt(490)
hero.fill.solid()
hero.fill.fore_color.rgb = C_CARD_CREAM
hero.line.color.rgb = C_CARD_BORDER
hero.line.width = Pt(1)

# 3. Badge 1: Project Pill (Rounded Rectangle 4)
badge1 = slide.shapes[3]
badge1.left = Pt(65)
badge1.top = Pt(42)
badge1.width = Pt(225)
badge1.height = Pt(24)
badge1.fill.solid()
badge1.fill.fore_color.rgb = C_JET_BLACK
badge1.line.color.rgb = C_JET_BLACK
tf = badge1.text_frame
tf.word_wrap = False
tf.margin_left = Pt(8)
tf.margin_right = Pt(8)
tf.margin_top = Pt(2)
tf.margin_bottom = Pt(2)
p = tf.paragraphs[0]
p.text = "MAJOR PROJECT EVALUATION • 2026"
p.font.name = FONT_NAME
p.font.size = Pt(8.5)
p.font.bold = True
p.font.color.rgb = C_CREAM_TEXT
p.alignment = PP_ALIGN.CENTER

# 4. Badge 2: Guide & Co-Guide Pill
guide_shape = None
for sh in slide.shapes:
    if sh.name == "GuideBadge":
        guide_shape = sh
        break
if not guide_shape:
    guide_shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Pt(545), Pt(42), Pt(350), Pt(24))
    guide_shape.name = "GuideBadge"

guide_shape.left = Pt(545)
guide_shape.top = Pt(42)
guide_shape.width = Pt(350)
guide_shape.height = Pt(24)
guide_shape.fill.solid()
guide_shape.fill.fore_color.rgb = C_WHITE
guide_shape.line.color.rgb = C_CARD_BORDER
guide_shape.line.width = Pt(1)
gtf = guide_shape.text_frame
gtf.word_wrap = False
gtf.margin_left = Pt(8)
gtf.margin_right = Pt(8)
gtf.margin_top = Pt(2)
gtf.margin_bottom = Pt(2)
gp = gtf.paragraphs[0]
gp.text = "GUIDE: Ms. Kamini   |   CO-GUIDE: Ms. Mekhala"
gp.font.name = FONT_NAME
gp.font.size = Pt(9.5)
gp.font.bold = True
gp.font.color.rgb = C_JET_BLACK
gp.alignment = PP_ALIGN.CENTER

# 5. Title Box (TextBox 5)
title_box = slide.shapes[4]
title_box.left = Pt(65)
title_box.top = Pt(72)
title_box.width = Pt(830)
title_box.height = Pt(78)
ttf = title_box.text_frame
ttf.word_wrap = True
ttf.margin_left = Pt(0)
ttf.margin_right = Pt(0)
ttf.margin_top = Pt(0)
ttf.margin_bottom = Pt(0)
ttf.clear()

p0 = ttf.paragraphs[0]
p0.text = "AcademicUniverse"
p0.font.name = FONT_NAME
p0.font.size = Pt(28)
p0.font.bold = True
p0.font.color.rgb = C_JET_BLACK
p0.space_after = Pt(2)

p1 = ttf.add_paragraph()
p1.text = "An AI Powered Educational Ecosystem for Student Growth and Academic Performance Analysis"
p1.font.name = FONT_NAME
p1.font.size = Pt(13)
p1.font.bold = True
p1.font.color.rgb = C_CHARCOAL
p1.space_after = Pt(2)

p2 = ttf.add_paragraph()
p2.text = "Department of Computer Science & Engineering • School of Engineering and Technology • Sharda University"
p2.font.name = FONT_NAME
p2.font.size = Pt(9.5)
p2.font.bold = False
p2.font.color.rgb = C_MUTED

# 6. Team Cards (3 Cards)
cards_data = [
    {
        "shape": slide.shapes[5],
        "role_badge": "RESOURCE COORDINATOR",
        "name": "1. Kushagra Singh Bhadauria",
        "sys_id": "2023361009",
        "roll_no": "2301010463",
        "email": "2023361009.kushagra@ug.sharda.ac.in",
        "phone": "+91 6395248403",
        "responsibilities": [
            "Resource Gathering & Support Lead",
            "Collecting Datasets, PDFs & Timetables",
            "Gathering Technical References & Docs",
            "Documentation & Integration Support",
            "Operational Backbone for Project Execution"
        ]
    },
    {
        "shape": slide.shapes[6],
        "role_badge": "SYSTEM ARCHITECT & DEV LEAD",
        "name": "2. Aashish Rajput",
        "sys_id": "2023329421",
        "roll_no": "2301010007",
        "email": "2023329421.aashish@ug.sharda.ac.in",
        "phone": "+91 9319977285",
        "responsibilities": [
            "Next.js 16 App Router & System Architecture",
            "Modular Express Microservices & MongoDB",
            "Playwright E-Zone Timetable Sync Engine",
            "Context-Injected Gemini 2.5 Campus Chatbot",
            "N-Way Overlap Planner & CI/CD Pipeline"
        ]
    },
    {
        "shape": slide.shapes[7],
        "role_badge": "RESEARCH & ANALYSIS LEAD",
        "name": "3. Avdesh Kumar Sah",
        "sys_id": "2023265132",
        "roll_no": "2301010220",
        "email": "2023265132.avdesh@ug.sharda.ac.in",
        "phone": "+91 8287054712",
        "responsibilities": [
            "Student Growth & Academic Performance Analysis",
            "Empirical Benchmark & Evaluation Design",
            "Academic Research Literature Review",
            "Tamper-Evident Record & Storage Testing",
            "Analytical Validation & Project Audits"
        ]
    }
]

left_positions = [Pt(65), Pt(355), Pt(645)]

for idx, c in enumerate(cards_data):
    sh = c["shape"]
    sh.left = left_positions[idx]
    sh.top = Pt(165)
    sh.width = Pt(250)
    sh.height = Pt(332)
    sh.fill.solid()
    sh.fill.fore_color.rgb = C_WHITE
    sh.line.color.rgb = C_CARD_BORDER
    sh.line.width = Pt(1)
    
    ctf = sh.text_frame
    ctf.word_wrap = True
    ctf.margin_left = Pt(14)
    ctf.margin_right = Pt(14)
    ctf.margin_top = Pt(14)
    ctf.margin_bottom = Pt(10)
    ctf.clear()
    
    # 1. Role Badge
    cp0 = ctf.paragraphs[0]
    cp0.text = c["role_badge"]
    cp0.font.name = FONT_NAME
    cp0.font.size = Pt(8.5)
    cp0.font.bold = True
    cp0.font.color.rgb = C_JET_BLACK
    cp0.alignment = PP_ALIGN.LEFT
    cp0.space_after = Pt(2)
    
    # 2. Member Name
    cp1 = ctf.add_paragraph()
    cp1.text = c["name"]
    cp1.font.name = FONT_NAME
    cp1.font.size = Pt(12)
    cp1.font.bold = True
    cp1.font.color.rgb = C_JET_BLACK
    cp1.alignment = PP_ALIGN.LEFT
    cp1.space_after = Pt(4)
    
    # 3. System ID & Roll
    cp2 = ctf.add_paragraph()
    cp2.text = f"Sys ID: {c['sys_id']}  |  Roll: {c['roll_no']}"
    cp2.font.name = FONT_NAME
    cp2.font.size = Pt(8.5)
    cp2.font.bold = False
    cp2.font.color.rgb = C_CHARCOAL
    cp2.alignment = PP_ALIGN.LEFT
    cp2.space_after = Pt(2)
    
    # 4. Email
    cp3 = ctf.add_paragraph()
    cp3.text = f"✉  {c['email']}"
    cp3.font.name = FONT_NAME
    cp3.font.size = Pt(7.5)
    cp3.font.bold = False
    cp3.font.color.rgb = C_MUTED
    cp3.alignment = PP_ALIGN.LEFT
    cp3.space_after = Pt(1)
    
    # 5. Phone
    cp4 = ctf.add_paragraph()
    cp4.text = f"✆  {c['phone']}"
    cp4.font.name = FONT_NAME
    cp4.font.size = Pt(8.5)
    cp4.font.bold = False
    cp4.font.color.rgb = C_MUTED
    cp4.alignment = PP_ALIGN.LEFT
    cp4.space_after = Pt(8)
    
    # 6. Section Header
    cp5 = ctf.add_paragraph()
    cp5.text = "CORE RESPONSIBILITIES"
    cp5.font.name = FONT_NAME
    cp5.font.size = Pt(8)
    cp5.font.bold = True
    cp5.font.color.rgb = C_JET_BLACK
    cp5.alignment = PP_ALIGN.LEFT
    cp5.space_after = Pt(4)
    
    # 7. Bullets
    for resp in c["responsibilities"]:
        bp = ctf.add_paragraph()
        bp.text = f"•  {resp}"
        bp.font.name = FONT_NAME
        bp.font.size = Pt(8)
        bp.font.bold = False
        bp.font.color.rgb = C_CHARCOAL
        bp.alignment = PP_ALIGN.LEFT
        bp.space_after = Pt(3)

# 7. Update Footers on Slides 2 to 10
for s_idx in range(1, 10):
    s = prs.slides[s_idx]
    for sh in s.shapes:
        if sh.has_text_frame:
            for p in sh.text_frame.paragraphs:
                if "Presenter: Aashish Rajput" in p.text or "AcademicUniverse | Guide" in p.text:
                    p.text = "AcademicUniverse | Guide: Ms. Kamini, Co-Guide: Ms. Mekhala | Team: Kushagra, Aashish, Avdesh"
                    p.font.name = FONT_NAME
                    p.font.color.rgb = C_MUTED
                    p.font.size = Pt(8.5)

prs.save("pptv2.pptx")
print("Saved pptv2.pptx successfully!")

# Export Slide 1 preview
abs_pptx = os.path.abspath("pptv2.pptx")
abs_png1 = os.path.abspath("scratch/final_slides/slide_1.png")
ppt = win32com.client.Dispatch('PowerPoint.Application')
pres = ppt.Presentations.Open(abs_pptx, WithWindow=False)
pres.Slides(1).Export(abs_png1, 'PNG', 1920, 1080)
pres.Close()
ppt.Quit()
print("Exported final slide 1 preview!")
