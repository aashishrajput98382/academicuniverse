import subprocess
import fitz  # PyMuPDF
import docx
import win32com.client
import pypdf
from docx.shared import Inches, Pt
from pathlib import Path

workspace = Path(__file__).resolve().parents[1]
svg_path = workspace / "docs" / "figures" / "au.drawio.svg"
fig1_png = workspace / "docs" / "paper" / "figure1_system_architecture_mermaid.png"
v39_docx = workspace / "docs" / "paper" / "PaperV39_Ollama_Primary.docx"
v40_docx = workspace / "docs" / "paper" / "PaperV40_Ollama_Primary.docx"
v40_pdf  = workspace / "docs" / "paper" / "PaperV40_Ollama_Primary.pdf"

# 1. Render SVG to high-res PNG (300 DPI)
print("[STEP 1] Rendering au.drawio.svg -> PNG at 300 DPI...")
doc_svg = fitz.open(str(svg_path.resolve()))
page = doc_svg[0]
pix = page.get_pixmap(dpi=300)
pix.save(str(fig1_png.resolve()))
print(f"  [SUCCESS] Rendered {fig1_png.name} ({pix.width}x{pix.height})")

# 2. Replace Fig. 1 image in DOCX
print("[STEP 2] Replacing Fig. 1 image in DOCX...")
doc = docx.Document(v39_docx)

for i, p in enumerate(doc.paragraphs):
    if p.text.strip().startswith("Fig. 1.") and "System Architecture" in p.text:
        img_p = doc.paragraphs[i - 1]
        img_p.clear()
        img_p.alignment = docx.enum.text.WD_ALIGN_PARAGRAPH.CENTER
        img_p.paragraph_format.space_before = Pt(2)
        img_p.paragraph_format.space_after = Pt(2)
        r = img_p.add_run()
        r.add_picture(str(fig1_png.resolve()), width=Inches(3.6))
        print(f"  [SUCCESS] Replaced Fig. 1 image at paragraph {i - 1}")
        break

doc.save(v40_docx)
print(f"  [SUCCESS] Saved {v40_docx.name}")

# 3. Export to PDF via Word COM
print("[STEP 3] Exporting to PDF via Word COM...")
subprocess.run(["taskkill", "/F", "/IM", "WINWORD.EXE"], capture_output=True)

word = win32com.client.DispatchEx("Word.Application")
word.Visible = False
word.DisplayAlerts = False

try:
    doc_com = word.Documents.Open(str(v40_docx.resolve()))
    doc_com.SaveAs(str(v40_pdf.resolve()), FileFormat=17)
    doc_com.Close(False)
    print(f"  [SUCCESS] Exported {v40_pdf.name}")
finally:
    word.Quit()
    subprocess.run(["taskkill", "/F", "/IM", "WINWORD.EXE"], capture_output=True)

# 4. Verify page count
reader = pypdf.PdfReader(str(v40_pdf))
pages = len(reader.pages)
print(f"\n[INFO] Paper V40 PDF Page Count: {pages} pages")
assert pages <= 20, f"Error: PDF exceeds 20 pages ({pages})!"
print(f"[PASS] V40 PDF within page budget: {pages} <= 20 pages")
