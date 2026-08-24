import os
import docx
import win32com.client
from docx.shared import Pt
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls
from pathlib import Path

workspace = Path(__file__).resolve().parents[1]
v23_docx_path = workspace / "docs" / "paper" / "PaperV23_Ollama_Primary.docx"
v23_pdf_path = workspace / "docs" / "paper" / "PaperV23_Ollama_Primary.pdf"
fig_dir = workspace / "docs" / "paper" / "extracted_figures"

doc = docx.Document(v23_docx_path)

caption_map = {
    "Fig. 1.": "image1.png",
    "Fig. 2.": "image2.png",
    "Fig. 3.": "image3.png",
    "Fig. 4.": "image4.png",
    "Fig. 5.": "image5.png",
    "Fig. 6.": "image6.png",
    "Fig. 7.": "image7.png",
    "Fig. 8.": "image8.png",
    "Fig. 9.": "image9.png",
}

print("=== EXACT MAPPING OF ALL 9 FIGURES TO THEIR CAPTIONS ===")

for prefix, img_name in caption_map.items():
    found = False
    for idx, p in enumerate(doc.paragraphs):
        txt = p.text.strip()
        if txt.startswith(prefix):
            prev_p = doc.paragraphs[idx-1]
            blips = prev_p._element.xpath('.//a:blip')
            for b in blips:
                embed_id = b.get('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed')
                rel = doc.part.rels[embed_id]
                img_path = fig_dir / img_name
                with open(img_path, "rb") as f:
                    rel.target_part._blob = f.read()
                print(f"[REPLACED] {prefix} (Caption P[{idx:03d}]) -> embed_id={embed_id} ({rel.target_ref}) with {img_name}")
                found = True
            break
    if not found:
        print(f"[WARNING] Could not find caption starting with '{prefix}'")

# Adjust Drawing Extents
for p in doc.paragraphs:
    for r in p.runs:
        drawings = r._r.xpath('.//w:drawing')
        if drawings:
            for d in drawings:
                exts = d.xpath('.//wp:extent')
                for ext in exts:
                    ext.set('cx', str(int(6.0 * 914400)))
                    ext.set('cy', str(int(3.2 * 914400)))
                a_exts = d.xpath('.//a:ext')
                for a_ext in a_exts:
                    a_ext.set('cx', str(int(6.0 * 914400)))
                    a_ext.set('cy', str(int(3.2 * 914400)))

doc.save(v23_docx_path)
print(f"[SUCCESS] Saved Word document!")

# Export PDF
word = None
try:
    word = win32com.client.DispatchEx("Word.Application")
    word.Visible = False
    word.DisplayAlerts = False
    doc_word = word.Documents.Open(os.path.abspath(str(v23_docx_path)), ReadOnly=True)
    doc_word.SaveAs(os.path.abspath(str(v23_pdf_path)), FileFormat=17)
    page_count = doc_word.ComputeStatistics(2)
    doc_word.Close(SaveChanges=False)
    print(f"[SUCCESS] Exported high-quality PDF: {v23_pdf_path.name} ({page_count} Pages!)")
except Exception as e:
    print(f"Word COM PDF Export Error: {e}")
finally:
    if word:
        try: word.Quit()
        except: pass
