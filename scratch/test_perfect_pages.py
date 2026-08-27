import subprocess
import win32com.client
import docx
import fitz
from docx.shared import Inches, Pt
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls
from pathlib import Path

workspace = Path.cwd()
v42_docx = workspace / "docs" / "paper" / "PaperV42_Ollama_Primary.docx"
fig1_png = workspace / "docs" / "paper" / "figure1_system_architecture_mermaid.png"
fig2_png = workspace / "docs" / "paper" / "figure2_data_flow_mermaid.png"
if not fig2_png.exists():
    fig2_png = workspace / "docs" / "paper" / "methodology_workflow_600dpi.png"

test_docx = workspace / "scratch" / "test_perfect_pages.docx"
test_pdf = workspace / "scratch" / "test_perfect_pages.pdf"

doc = docx.Document(v42_docx)

# Resize Fig 1 and Fig 2 so image + caption + explanation fit together on their own single page
for i, p in enumerate(doc.paragraphs):
    if p.text.strip().startswith("Fig. 1.") and "System Architecture" in p.text:
        img_p = doc.paragraphs[i - 1]
        img_p.clear()
        img_p.alignment = docx.enum.text.WD_ALIGN_PARAGRAPH.CENTER
        img_p.paragraph_format.space_before = Pt(2)
        img_p.paragraph_format.space_after = Pt(2)
        img_p.paragraph_format.keep_with_next = True
        r = img_p.add_run()
        r.add_picture(str(fig1_png.resolve()), width=Inches(3.1))
        
        p.paragraph_format.keep_with_next = True
        p.paragraph_format.keep_lines = True
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(3)

    if p.text.strip().startswith("Fig. 2.") and "Data Flow" in p.text:
        img_p = doc.paragraphs[i - 1]
        img_p.clear()
        img_p.alignment = docx.enum.text.WD_ALIGN_PARAGRAPH.CENTER
        img_p.paragraph_format.space_before = Pt(2)
        img_p.paragraph_format.space_after = Pt(2)
        img_p.paragraph_format.keep_with_next = True
        r = img_p.add_run()
        r.add_picture(str(fig2_png.resolve()), width=Inches(2.7))
        
        p.paragraph_format.keep_with_next = True
        p.paragraph_format.keep_lines = True
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(3)

# Glues all other images and tables to their captions
for i, p in enumerate(doc.paragraphs):
    has_img = len(p._p.xpath('.//w:drawing')) > 0
    is_tbl_cap = p.text.strip().startswith('TABLE ')
    is_fig_cap = p.text.strip().startswith('Fig. ') or p.text.strip().startswith('Figure ')
    
    if has_img:
        p.paragraph_format.keep_with_next = True
        p.paragraph_format.keep_lines = True
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(2)
    if is_tbl_cap:
        p.paragraph_format.keep_with_next = True
        p.paragraph_format.keep_lines = True
        p.paragraph_format.space_before = Pt(3)
        p.paragraph_format.space_after = Pt(1.5)
    if is_fig_cap:
        p.paragraph_format.keep_lines = True
        p.paragraph_format.space_before = Pt(1.5)
        p.paragraph_format.space_after = Pt(3)

for tbl in doc.tables:
    for row in tbl.rows:
        trPr = row._tr.get_or_add_trPr()
        trPr.append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>'))

doc.save(test_docx)

# Export to PDF
subprocess.run(["taskkill", "/F", "/IM", "WINWORD.EXE"], capture_output=True)
word = win32com.client.DispatchEx("Word.Application")
word.Visible = False
word.DisplayAlerts = False
try:
    d = word.Documents.Open(str(test_docx.resolve()))
    d.SaveAs(str(test_pdf.resolve()), FileFormat=17)
    d.Close(False)
finally:
    word.Quit()
    subprocess.run(["taskkill", "/F", "/IM", "WINWORD.EXE"], capture_output=True)

# Inspect pages
pdf_doc = fitz.open(str(test_pdf))
print(f"Total PDF Pages: {len(pdf_doc)}")
for page_num in range(len(pdf_doc)):
    page = pdf_doc[page_num]
    images = page.get_images()
    text = page.get_text()
    captions = [l.strip() for l in text.split('\n') if l.strip().startswith('Fig. ') or l.strip().startswith('Figure ')]
    tbl_captions = [l.strip() for l in text.split('\n') if l.strip().startswith('TABLE ')]
    if images or captions or tbl_captions:
        print(f"Page {page_num+1:2d}: Images={len(images)} | Fig Caps={len(captions)} | Tbl Caps={len(tbl_captions)}")
        for c in captions: print(f"    [FIG CAP] {c[:70]}")
        for t in tbl_captions: print(f"    [TBL CAP] {t[:70]}")
