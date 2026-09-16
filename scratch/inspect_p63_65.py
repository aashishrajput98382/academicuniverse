import docx

doc = docx.Document('PaperV50.docx')
for idx in [62, 63, 64, 65, 66]:
    p = doc.paragraphs[idx]
    print(f'P{idx}: "{p.text}"')

t7 = doc.tables[7]
print(f'Table 7 rows={len(t7.rows)}, cols={len(t7.columns)}')
for r_idx in range(min(3, len(t7.rows))):
    row = t7.rows[r_idx]
    print(f'Row {r_idx}: {[c.text.strip() for c in row.cells]}')
