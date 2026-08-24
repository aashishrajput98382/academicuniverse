import os
import re
import subprocess
import win32com.client
import docx
from docx.shared import Pt, RGBColor
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls
from pathlib import Path

workspace = Path(__file__).resolve().parents[1]
v26_docx_path = workspace / "docs" / "paper" / "PaperV26_Ollama_Primary.docx"
v26_pdf_path = workspace / "docs" / "paper" / "PaperV26_Ollama_Primary.pdf"
fig_dir = workspace / "docs" / "paper" / "extracted_figures"

# 1. Run pipeline
subprocess.run(["python", "scratch/generate_paperv26_pipeline.py"], check=True)

# 2. Open docx
doc = docx.Document(v26_docx_path)

# 3. Apply Elegant IEEE Table Styling to ALL Tables
def set_cell_shading(cell, color_hex):
    shading_elm = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{color_hex}"/>')
    cell._tc.get_or_add_tcPr().append(shading_elm)

for tbl_idx, tbl in enumerate(doc.tables):
    for r_idx, row in enumerate(tbl.rows):
        is_header = (r_idx == 0)
        bg_color = "1B365D" if is_header else ("F7FAFC" if r_idx % 2 == 1 else "FFFFFF")
        
        for cell in row.cells:
            tcPr = cell._tc.get_or_add_tcPr()
            tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="28" w:type="dxa"/><w:bottom w:w="28" w:type="dxa"/><w:left w:w="35" w:type="dxa"/><w:right w:w="35" w:type="dxa"/></w:tcMar>')
            tcPr.append(tcMar)
            set_cell_shading(cell, bg_color)
            
            for p in cell.paragraphs:
                p.paragraph_format.space_before = Pt(1)
                p.paragraph_format.space_after = Pt(1)
                p.paragraph_format.line_spacing = 1.0
                for r in p.runs:
                    r.font.name = "Times New Roman"
                    if is_header:
                        r.font.size = Pt(8.0)
                        r.font.bold = True
                        r.font.color.rgb = RGBColor(255, 255, 255)
                    else:
                        r.font.size = Pt(7.2)
                        r.font.color.rgb = RGBColor(30, 41, 59)

# 4. Fix exact figure mappings for all 9 figures
caption_map = {
    "Fig. 1.": "image1.png",
    "Fig. 2.": "image2.png",
    "Fig. 3.": "image3.png",
    "Fig. 4.": "image4.png",
    "Fig. 5.": "image5.png",
    "Fig. 6.": "image6.png",
    "Fig. 7.": "image7.png",
    "Fig. 8.": "image8.png",
    "Fig. 9.": "image9.png",
}

print("=== EXACT MAPPING OF ALL 9 FIGURES IN V26 ===")
for prefix, img_name in caption_map.items():
    for idx, p in enumerate(doc.paragraphs):
        txt = p.text.strip()
        if txt.startswith(prefix):
            prev_p = doc.paragraphs[idx-1]
            blips = prev_p._element.xpath('.//a:blip')
            for b in blips:
                embed_id = b.get('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed')
                rel = doc.part.rels[embed_id]
                img_path = fig_dir / img_name
                with open(img_path, "rb") as f:
                    rel.target_part._blob = f.read()
                print(f"[REPLACED] {prefix} -> {img_name}")
            break

# 5. Compact paragraph line spacing and drawing dimensions for <= 20 pages
for p in doc.paragraphs:
    txt = p.text.strip()
    if txt.startswith("TABLE ") or txt.startswith("Table "):
        p.paragraph_format.space_before = Pt(3)
        p.paragraph_format.space_after = Pt(1)
    elif txt.startswith("Fig. ") or txt.startswith("FIG. "):
        p.paragraph_format.space_before = Pt(1)
        p.paragraph_format.space_after = Pt(3)
    elif re.match(r'^\d+\.\s+', txt):
        p.paragraph_format.space_before = Pt(6)
        p.paragraph_format.space_after = Pt(2)
    else:
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.line_spacing = 1.05

    for r in p.runs:
        drawings = r._r.xpath('.//w:drawing')
        if drawings:
            for d in drawings:
                exts = d.xpath('.//wp:extent')
                for ext in exts:
                    ext.set('cx', str(int(6.0 * 914400)))
                    ext.set('cy', str(int(3.2 * 914400)))
                a_exts = d.xpath('.//a:ext')
                for a_ext in a_exts:
                    a_ext.set('cx', str(int(6.0 * 914400)))
                    a_ext.set('cy', str(int(3.2 * 914400)))

doc.save(v26_docx_path)
print(f"[SUCCESS] Saved PaperV26_Ollama_Primary.docx!")

# 6. Export PDF
subprocess.run(["taskkill", "/F", "/IM", "WINWORD.EXE"], capture_output=True)

word = win32com.client.DispatchEx("Word.Application")
word.Visible = False
word.DisplayAlerts = False
try:
    doc_word = word.Documents.Open(os.path.abspath(str(v26_docx_path)))
    doc_word.SaveAs(os.path.abspath(str(v26_pdf_path)), FileFormat=17)
    page_count = doc_word.ComputeStatistics(2)
    doc_word.Close(False)
    print(f"[SUCCESS] Exported high-quality PDF: {v26_pdf_path.name} ({page_count} Pages!)")
finally:
    word.Quit()
