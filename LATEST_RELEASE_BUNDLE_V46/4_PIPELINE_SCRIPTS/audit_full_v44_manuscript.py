import docx
import re
import fitz

doc = docx.Document('docs/paper/PaperV44_Ollama_Primary.docx')

print("=================================================================")
print(" COMPREHENSIVE FORENSIC AUDIT OF PAPER V44 MANUSCRIPT")
print("=================================================================")

# 1. Check for any leftover obsolete numbers
obsolete_patterns = [
    (r'\b45\s+(?:synthetic|physical|document|specimen)', 'Old 45 document count'),
    (r'\b900\s+(?:paired|field|atomic|evaluated)', 'Old 900 field count'),
    (r'\b97\.01\b', 'Old McNemar statistic 97.01'),
    (r'\b79\.56%', 'Old Pass A F1 79.56%'),
    (r'\b90\.56%', 'Old Pass B F1 90.56%'),
    (r'\b9\.69%', 'Old Pass A CER 9.69%'),
    (r'\b3\.64%', 'Old Pass B CER 3.64%'),
    (r'\b62\.44%', 'Old relative CER reduction 62.44%'),
    (r'\b11\.00%', 'Old F1 gain 11.00%'),
    (r'\b2917\.90', 'Old runtime duration 2917.90s'),
    (r'\b48\.63\s+mins', 'Old runtime duration 48.63 mins'),
    (r'\brun_1785959173886\b', 'Old run ID'),
    (r'\b0bc63b0\b', 'Old commit hash 0bc63b0')
]

obsolete_findings = []

for i, p in enumerate(doc.paragraphs):
    for pat, desc in obsolete_patterns:
        matches = re.findall(pat, p.text, re.IGNORECASE)
        if matches:
            obsolete_findings.append(f"Paragraph {i}: Found '{matches}' ({desc}) -> \"{p.text[:80]}...\"")

for t_idx, tbl in enumerate(doc.tables):
    for r_idx, row in enumerate(tbl.rows):
        for c_idx, cell in enumerate(row.cells):
            for pat, desc in obsolete_patterns:
                matches = re.findall(pat, cell.text, re.IGNORECASE)
                if matches:
                    obsolete_findings.append(f"Table {t_idx+1} R{r_idx} C{c_idx}: Found '{matches}' ({desc}) -> \"{cell.text.strip()[:60]}\"")

print(f"\n[AUDIT CHECK 1: OBSOLETE NUMBERS & STRINGS]")
if not obsolete_findings:
    print("  >>> PASS: ZERO obsolete numbers or strings found in entire DOCX!")
else:
    print(f"  >>> WARNING: Found {len(obsolete_findings)} occurrences:")
    for f in obsolete_findings: print("     ", f)

# 2. Check Table Inventory & Integrity
print(f"\n[AUDIT CHECK 2: TABLE INVENTORY & INTEGRITY]")
print(f"  Total Tables in DOCX: {len(doc.tables)}")
assert len(doc.tables) == 14, f"Error: Expected 14 tables, found {len(doc.tables)}!"

table_names = [
    "Table 1: Literature Survey",
    "Table 2: Computing Environment",
    "Table 3: Multi-Modal Dataset",
    "Table 4: Optical Degradation Profiles",
    "Table 5: Canonical Parameters",
    "Table 6: Mathematical Metrics",
    "Table 7: Dry-Run Verification",
    "Table 8: SOTA Empirical Benchmark",
    "Table 9: Normalization Ablation",
    "Table 10: Rule Corrections",
    "Table 11: Statistical Hypothesis Tests",
    "Table 12: 10k Bootstrap CIs",
    "Table 13: 9-Class Error Taxonomy",
    "Table 14: ML Failure Prediction"
]

for idx, tbl in enumerate(doc.tables):
    header = [c.text.strip().replace('\n', ' ') for c in tbl.rows[0].cells]
    print(f"  -> Table {idx+1:2d} ({len(tbl.rows):2d} rows x {len(tbl.columns):2d} cols): {table_names[idx]} | Header: {header[:2]}")

# 3. Check PDF Pages and Figure/Table Co-location
print(f"\n[AUDIT CHECK 3: PDF PAGE COUNT & FIGURE/TABLE CO-LOCATION]")
pdf = fitz.open('docs/paper/PaperV44_Ollama_Primary.pdf')
print(f"  Total PDF Pages: {len(pdf)}")
assert len(pdf) == 20, f"Error: PDF page count {len(pdf)} is not 20!"

for p_num in range(len(pdf)):
    page = pdf[p_num]
    imgs = page.get_images()
    txt = page.get_text()
    figs = [l.strip() for l in txt.split('\n') if l.strip().startswith('Fig. ') or l.strip().startswith('Figure ')]
    tbls = [l.strip() for l in txt.split('\n') if l.strip().startswith('TABLE ')]
    if imgs or figs or tbls:
        print(f"  Page {p_num+1:2d}: Images={len(imgs)} | Figures={len(figs)} | Tables={len(tbls)}")
        for f in figs: print(f"       [FIG] {f[:65]}")
        for t in tbls: print(f"       [TBL] {t[:65]}")

print("\n=================================================================")
print(" ALL AUDIT CHECKS COMPLETED!")
print("=================================================================")
